"""冻存管与冻存盒服务。"""

from __future__ import annotations

import sqlite3
from datetime import datetime
from typing import Any, BinaryIO

from openpyxl import load_workbook

from ..domain import (
    VIAL_STATUSES,
    normalize_choice,
    normalize_date,
    normalize_int,
    normalize_number,
    normalize_text,
    parse_cell_text,
    today_iso,
)
from ..errors import Conflict, ValidationError
from . import audit, fetch_all, fetch_one, get_or_404

GRID_ROWS = "ABCDEFGHI"
GRID_COLS = range(1, 10)
GRID_CAPACITY = len(GRID_ROWS) * len(GRID_COLS)


def split_location(location: str | None) -> tuple[str | None, str | None]:
    """把「罐A/架1/盒1/A1」拆成 (盒路径, 孔位)。

    兼容两种写法：3 段盒路径 + 孔位（旧版），或 2 段（盒名/孔位，例如 Excel 导入）。
    """
    text = (location or "").strip()
    if not text:
        return None, None
    parts = [part.strip() for part in text.split("/") if part.strip()]
    if len(parts) < 2:
        return None, None
    box = "/".join(parts[:-1])
    hole = parts[-1]
    if hole.startswith("孔"):
        hole = hole[1:]
    return box, hole.upper()


def _hole_number(hole: str) -> tuple[str, int] | None:
    hole = hole.strip().upper()
    if len(hole) < 2:
        return None
    row, col = hole[0], hole[1:]
    if row not in GRID_ROWS or not col.isdigit():
        return None
    index = int(col)
    if index not in GRID_COLS:
        return None
    return row, index


def list_vials(
    conn: sqlite3.Connection,
    q: str | None = None,
    cell_line_id: int | None = None,
    status: str | None = None,
    limit: int = 500,
    unplaced: bool = False,
) -> list[dict[str, Any]]:
    sql = (
        "SELECT fv.*, cl.name AS cell_name, cb.batch_name "
        "FROM frozen_vials fv LEFT JOIN cell_lines cl ON cl.id = fv.cell_line_id "
        "LEFT JOIN culture_batches cb ON cb.id = fv.batch_id WHERE 1 = 1"
    )
    params: list[Any] = []
    if q:
        sql += " AND (cl.name LIKE ? OR fv.location LIKE ? OR fv.notes LIKE ?)"
        params.extend([f"%{q.strip()}%"] * 3)
    if cell_line_id:
        sql += " AND fv.cell_line_id = ?"
        params.append(cell_line_id)
    if status:
        sql += " AND fv.status = ?"
        params.append(status)
    if unplaced:
        sql += " AND (fv.location IS NULL OR TRIM(fv.location) = '')"
    sql += " ORDER BY fv.location, fv.id LIMIT ?"
    params.append(max(1, min(int(limit), 2000)))
    return fetch_all(conn, sql, params)


def get_vial(conn: sqlite3.Connection, vial_id: int) -> dict[str, Any]:
    row = fetch_one(
        conn,
        "SELECT fv.*, cl.name AS cell_name FROM frozen_vials fv "
        "LEFT JOIN cell_lines cl ON cl.id = fv.cell_line_id WHERE fv.id = ?",
        (vial_id,),
    )
    if row is None:
        from ..errors import NotFound

        raise NotFound(f"冻存管不存在（id={vial_id}）")
    return row


def _resolve_line(conn: sqlite3.Connection, payload: dict[str, Any]) -> dict[str, Any]:
    cell_line_id = normalize_int(payload.get("cell_line_id"), "细胞株", required=True)
    return get_or_404(conn, "cell_lines", int(cell_line_id), "细胞株")


def create_vials(conn: sqlite3.Connection, payload: dict[str, Any],
                 user: str | None = None) -> list[dict[str, Any]]:
    line = _resolve_line(conn, payload)
    freeze_date = normalize_date(payload.get("freeze_date"), "冻存日期") or today_iso()
    cell_count = normalize_number(payload.get("cell_count"), "每管细胞数", minimum=0)
    pd_value = normalize_number(payload.get("pd"), "冻存时PD值")
    status = normalize_choice(payload.get("status"), VIAL_STATUSES, "状态", default="在库")
    notes = normalize_text(payload.get("notes"))
    count = normalize_int(payload.get("count"), "添加管数", minimum=1) or 1

    raw_location = str(payload.get("location") or "").strip()
    locations = [part.strip() for part in raw_location.replace("，", ",").split(",") if part.strip()]
    if not locations:
        locations = [""]

    if count * len(locations) > 200:
        raise ValidationError("一次最多添加 200 支冻存管")

    created_ids: list[int] = []
    for location in locations:
        for _ in range(count):
            cursor = conn.execute(
                "INSERT INTO frozen_vials (cell_line_id, freeze_date, cell_count, location, pd, status, notes) "
                "VALUES (?,?,?,?,?,?,?)",
                (line["id"], freeze_date, cell_count, location, pd_value, status, notes),
            )
            created_ids.append(int(cursor.lastrowid))
    audit(conn, user, "create", "vial", created_ids[0] if created_ids else None,
          {"count": len(created_ids), "location": raw_location})
    return [get_vial(conn, vial_id) for vial_id in created_ids]


def update_vial(conn: sqlite3.Connection, vial_id: int, payload: dict[str, Any],
                user: str | None = None) -> dict[str, Any]:
    get_or_404(conn, "frozen_vials", vial_id, "冻存管")
    assignments: dict[str, Any] = {}
    if "cell_line_id" in payload:
        line = _resolve_line(conn, payload)
        assignments["cell_line_id"] = line["id"]
    if "location" in payload:
        assignments["location"] = normalize_text(payload.get("location"), 200)
    if "cell_count" in payload:
        assignments["cell_count"] = normalize_number(payload.get("cell_count"), "细胞数", minimum=0)
    if "pd" in payload:
        assignments["pd"] = normalize_number(payload.get("pd"), "PD值")
    if "freeze_date" in payload:
        assignments["freeze_date"] = normalize_date(payload.get("freeze_date"), "冻存日期")
    if "status" in payload:
        assignments["status"] = normalize_choice(payload.get("status"), VIAL_STATUSES, "状态")
    if "notes" in payload:
        assignments["notes"] = normalize_text(payload.get("notes"))
    if "batch_id" in payload:
        assignments["batch_id"] = normalize_int(payload.get("batch_id"), "关联批次")

    if assignments:
        clause = ", ".join(f"{key} = ?" for key in assignments)
        conn.execute(f"UPDATE frozen_vials SET {clause} WHERE id = ?", (*assignments.values(), vial_id))
        audit(conn, user, "update", "vial", vial_id, {"fields": list(assignments)})
    return get_vial(conn, vial_id)


def delete_vial(conn: sqlite3.Connection, vial_id: int, user: str | None = None) -> dict[str, Any]:
    vial = get_or_404(conn, "frozen_vials", vial_id, "冻存管")
    conn.execute("DELETE FROM frozen_vials WHERE id = ?", (vial_id,))
    audit(conn, user, "delete", "vial", vial_id, {"location": vial["location"]})
    return {"deleted": vial_id}


def list_boxes(conn: sqlite3.Connection) -> list[dict[str, Any]]:
    rows = conn.execute(
        "SELECT location FROM frozen_vials WHERE location IS NOT NULL AND location <> ''"
    ).fetchall()
    boxes: dict[str, int] = {}
    for row in rows:
        box, _ = split_location(row["location"])
        if box:
            boxes[box] = boxes.get(box, 0) + 1
    return [
        {"box": box, "occupied": count, "capacity": GRID_CAPACITY, "free": GRID_CAPACITY - count}
        for box, count in sorted(boxes.items())
    ]


def box_detail(conn: sqlite3.Connection, box: str) -> dict[str, Any]:
    vials = fetch_all(
        conn,
        "SELECT fv.*, cl.name AS cell_name FROM frozen_vials fv "
        "LEFT JOIN cell_lines cl ON cl.id = fv.cell_line_id "
        "WHERE fv.location LIKE ? ORDER BY fv.location",
        (f"{box}/%",),
    )
    holes: dict[str, Any] = {}
    unplaced: list[dict[str, Any]] = []
    for vial in vials:
        _, hole = split_location(vial["location"])
        if hole and _hole_number(hole):
            holes[hole] = vial
        else:
            unplaced.append(vial)
    return {
        "box": box,
        "rows": list(GRID_ROWS),
        "cols": list(GRID_COLS),
        "capacity": GRID_CAPACITY,
        "occupied": len(holes),
        "holes": holes,
        "unplaced": unplaced,
    }


def import_from_workbook(
    conn: sqlite3.Connection,
    stream: BinaryIO,
    *,
    mode: str = "append",
    user: str | None = None,
) -> dict[str, Any]:
    """按旧版约定导入：每个 Sheet 是一个冻存盒，A1:I9 为 9x9 孔位。"""
    if mode not in ("append", "replace"):
        raise ValidationError("导入模式必须是 append 或 replace")
    try:
        workbook = load_workbook(stream, data_only=True)
    except Exception as exc:  # noqa: BLE001 - openpyxl 的异常类型较多
        raise ValidationError(f"无法读取 Excel 文件：{exc}") from exc

    inserted = 0
    skipped = 0
    errors: list[str] = []
    lines_created: list[str] = []
    box_stats: list[dict[str, Any]] = []
    today = datetime.now().strftime("%Y-%m-%d")

    for sheet_name in workbook.sheetnames:
        box = str(sheet_name).strip()
        if not box:
            continue
        if mode == "replace":
            conn.execute("DELETE FROM frozen_vials WHERE location LIKE ?", (f"{box}/%",))

        box_inserted = 0
        box_skipped = 0
        sheet = workbook[sheet_name]
        for row_index in range(1, 10):
            row_letter = chr(64 + row_index)
            for col_index in range(1, 10):
                cell = sheet.cell(row=row_index, column=col_index)
                parsed = parse_cell_text(cell.value)
                if not parsed:
                    continue
                hole = f"{row_letter}{col_index}"
                location = f"{box}/{hole}"

                if mode == "append":
                    exists = conn.execute(
                        "SELECT id FROM frozen_vials WHERE location = ? LIMIT 1", (location,)
                    ).fetchone()
                    if exists:
                        box_skipped += 1
                        continue

                line = conn.execute(
                    "SELECT id, name FROM cell_lines WHERE name = ?", (parsed["cell_name"],)
                ).fetchone()
                if line:
                    line_id = int(line["id"])
                else:
                    cursor = conn.execute(
                        "INSERT INTO cell_lines (name, type, revival_date) VALUES (?, '永生化', ?)",
                        (parsed["cell_name"], today),
                    )
                    line_id = int(cursor.lastrowid)
                    lines_created.append(parsed["cell_name"])

                try:
                    conn.execute(
                        "INSERT INTO frozen_vials (cell_line_id, location, freeze_date, cell_count, pd, status, notes) "
                        "VALUES (?,?,?,?,?,?,?)",
                        (line_id, location, parsed["freeze_date"], parsed["cell_count"],
                         parsed["pd"], parsed["status"], parsed["notes"]),
                    )
                    box_inserted += 1
                except sqlite3.Error as exc:
                    errors.append(f"{location}: {exc}")

        inserted += box_inserted
        skipped += box_skipped
        box_stats.append({"box": box, "inserted": box_inserted, "skipped": box_skipped})

    audit(conn, user, "import", "vial", None,
          {"inserted": inserted, "skipped": skipped, "boxes": len(box_stats)})
    return {
        "mode": mode,
        "inserted": inserted,
        "skipped": skipped,
        "errors": errors[:50],
        "boxes": box_stats,
        "lines_created": lines_created,
    }


def clear_box(conn: sqlite3.Connection, box: str, user: str | None = None) -> dict[str, Any]:
    if not box.strip():
        raise ValidationError("盒子名称不能为空")
    deleted = conn.execute("DELETE FROM frozen_vials WHERE location LIKE ?", (f"{box}/%",)).rowcount
    audit(conn, user, "delete_box", "vial", None, {"box": box, "deleted": deleted})
    return {"box": box, "deleted": deleted}


def require_existing_box(conn: sqlite3.Connection, box: str) -> None:
    exists = conn.execute(
        "SELECT 1 FROM frozen_vials WHERE location LIKE ? LIMIT 1", (f"{box}/%",)
    ).fetchone()
    if not exists:
        raise Conflict(f"冻存盒「{box}」不存在或没有冻存管")
