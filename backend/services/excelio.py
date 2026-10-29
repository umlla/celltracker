"""Excel / CSV 导出（列名与旧版导出一致）。"""

from __future__ import annotations

import csv
import io
import sqlite3
from datetime import date
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import fetch_all, placeholders
from .groups import get_group
from ..errors import NotFound

HEADER_FILL = PatternFill("solid", fgColor="E8F1FB")
HEADER_FONT = Font(bold=True)


def _autosize(worksheet) -> None:
    for column_cells in worksheet.columns:
        longest = 8
        letter = None
        for cell in column_cells:
            if letter is None:
                letter = cell.column_letter
            value = "" if cell.value is None else str(cell.value)
            width = sum(2 if ord(char) > 127 else 1 for char in value)
            longest = max(longest, min(width + 2, 48))
        if letter:
            worksheet.column_dimensions[letter].width = longest


def _write_sheet(workbook: Workbook, title: str, headers: list[str],
                 rows: list[list[Any]], first: bool = False) -> None:
    worksheet = workbook.active if first else workbook.create_sheet()
    worksheet.title = title[:31]
    worksheet.append(headers)
    for cell in worksheet[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in rows:
        worksheet.append(["" if value is None else value for value in row])
    worksheet.freeze_panes = "A2"
    _autosize(worksheet)


def _rows(conn: sqlite3.Connection, sql: str, params: tuple[Any, ...] = ()) -> list[list[Any]]:
    return [list(row) for row in conn.execute(sql, params).fetchall()]


def group_export(conn: sqlite3.Connection, group_id: int, include_all: bool = True) -> io.BytesIO:
    group = get_group(conn, group_id)
    member_ids = [int(member["cell_line_id"]) for member in group["members"]]
    if not member_ids:
        raise NotFound("该细胞组没有成员，无法导出")
    marks = placeholders(member_ids)

    lines = _rows(
        conn,
        f"""
        SELECT id AS ID, name AS 细胞名称, type AS 类型, revival_date AS 复苏日期,
               default_passage_interval AS 默认传代间隔, notes AS 备注
        FROM cell_lines WHERE id IN ({marks}) ORDER BY name
        """,
        tuple(member_ids),
    )

    batch_filter = "" if include_all else " AND cb.is_active = 1"
    batches = _rows(
        conn,
        f"""
        SELECT cb.id AS ID, cl.name AS 细胞名称, cb.batch_name AS 批次名称,
               cb.revival_date AS 复苏日期, cb.source_freeze_date AS 冻存来源日期,
               cb.initial_pd AS 初始PD, cb.initial_cell_count AS 初始细胞数,
               CASE cb.is_active WHEN 1 THEN '是' ELSE '否' END AS 是否活跃,
               cb.end_date AS 终止日期, cb.end_reason AS 终止原因, cb.notes AS 备注
        FROM culture_batches cb JOIN cell_lines cl ON cl.id = cb.cell_line_id
        WHERE cl.id IN ({marks}){batch_filter}
        ORDER BY cl.name, cb.revival_date, cb.id
        """,
        tuple(member_ids),
    )

    batch_marks = placeholders(member_ids)
    record_filter = "" if include_all else " AND cr.is_active = 1"
    records = _rows(
        conn,
        f"""
        SELECT cr.id AS ID, cl.name AS 细胞名称, cb.batch_name AS 批次名称,
               cr.passage AS 代次, cr.date AS 日期,
               cr.inoculum_cells AS 接种细胞数, cr.inoculum_dishes AS 盘数,
               cr.harvested_cells AS 收获细胞数, cr.pd AS 本次PD,
               cr.cumulative_pd AS 累计PD, cr.status AS 状态,
               cr.next_passage_date AS 建议下次传代, cr.notes AS 备注
        FROM culture_records cr
        JOIN culture_batches cb ON cb.id = cr.batch_id
        JOIN cell_lines cl ON cl.id = cb.cell_line_id
        WHERE cl.id IN ({batch_marks}){record_filter}
        ORDER BY cl.name, cb.batch_name, cr.date, cr.id
        """,
        tuple(member_ids),
    )

    vials = _rows(
        conn,
        f"""
        SELECT fv.id AS ID, cl.name AS 细胞名称, fv.freeze_date AS 冻存日期,
               fv.cell_count AS 细胞数, fv.location AS 位置, fv.pd AS PD,
               fv.status AS 状态, fv.notes AS 备注
        FROM frozen_vials fv JOIN cell_lines cl ON cl.id = fv.cell_line_id
        WHERE cl.id IN ({marks})
        ORDER BY cl.name, fv.location
        """,
        tuple(member_ids),
    )

    workbook = Workbook()
    _write_sheet(workbook, "细胞株信息", ["ID", "细胞名称", "类型", "复苏日期", "默认传代间隔", "备注"],
                 lines, first=True)
    _write_sheet(workbook, "培养批次",
                 ["ID", "细胞名称", "批次名称", "复苏日期", "冻存来源日期", "初始PD", "初始细胞数",
                  "是否活跃", "终止日期", "终止原因", "备注"], batches)
    _write_sheet(workbook, "传代记录",
                 ["ID", "细胞名称", "批次名称", "代次", "日期", "接种细胞数", "盘数", "收获细胞数",
                  "本次PD", "累计PD", "状态", "建议下次传代", "备注"], records)
    _write_sheet(workbook, "冻存管",
                 ["ID", "细胞名称", "冻存日期", "细胞数", "位置", "PD", "状态", "备注"], vials)

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def full_export(conn: sqlite3.Connection) -> io.BytesIO:
    lines = _rows(
        conn,
        "SELECT id AS ID, name AS 细胞名称, type AS 类型, revival_date AS 复苏日期, "
        "default_passage_interval AS 默认传代间隔, notes AS 备注 FROM cell_lines ORDER BY name",
    )
    batches = _rows(
        conn,
        "SELECT cb.id AS ID, cl.name AS 细胞名称, cb.batch_name AS 批次名称, cb.revival_date AS 复苏日期, "
        "cb.source_freeze_date AS 冻存来源日期, cb.initial_pd AS 初始PD, cb.initial_cell_count AS 初始细胞数, "
        "CASE cb.is_active WHEN 1 THEN '是' ELSE '否' END AS 是否活跃, cb.end_date AS 终止日期, "
        "cb.end_reason AS 终止原因, cb.notes AS 备注 FROM culture_batches cb "
        "JOIN cell_lines cl ON cl.id = cb.cell_line_id ORDER BY cl.name, cb.revival_date, cb.id",
    )
    records = _rows(
        conn,
        "SELECT cr.id AS ID, cl.name AS 细胞名称, cb.batch_name AS 批次名称, cr.passage AS 代次, "
        "cr.date AS 日期, cr.inoculum_cells AS 接种细胞数, cr.inoculum_dishes AS 盘数, "
        "cr.harvested_cells AS 收获细胞数, cr.pd AS 本次PD, cr.cumulative_pd AS 累计PD, "
        "cr.status AS 状态, cr.next_passage_date AS 建议下次传代, cr.notes AS 备注 "
        "FROM culture_records cr JOIN culture_batches cb ON cb.id = cr.batch_id "
        "JOIN cell_lines cl ON cl.id = cb.cell_line_id ORDER BY cl.name, cb.batch_name, cr.date, cr.id",
    )
    vials = _rows(
        conn,
        "SELECT fv.id AS ID, cl.name AS 细胞名称, fv.freeze_date AS 冻存日期, fv.cell_count AS 细胞数, "
        "fv.location AS 位置, fv.pd AS PD, fv.status AS 状态, fv.notes AS 备注 "
        "FROM frozen_vials fv JOIN cell_lines cl ON cl.id = fv.cell_line_id ORDER BY cl.name, fv.location",
    )
    groups = _rows(conn, "SELECT id AS ID, group_name AS 组名, notes AS 备注 FROM cell_groups ORDER BY group_name")
    members = _rows(
        conn,
        "SELECT gm.group_id AS 组ID, g.group_name AS 组名, cl.name AS 细胞名称, gm.sort_order AS 排序 "
        "FROM cell_group_members gm JOIN cell_groups g ON g.id = gm.group_id "
        "JOIN cell_lines cl ON cl.id = gm.cell_line_id ORDER BY g.group_name, gm.sort_order",
    )

    workbook = Workbook()
    _write_sheet(workbook, "细胞株信息", ["ID", "细胞名称", "类型", "复苏日期", "默认传代间隔", "备注"],
                 lines, first=True)
    _write_sheet(workbook, "培养批次",
                 ["ID", "细胞名称", "批次名称", "复苏日期", "冻存来源日期", "初始PD", "初始细胞数",
                  "是否活跃", "终止日期", "终止原因", "备注"], batches)
    _write_sheet(workbook, "传代记录",
                 ["ID", "细胞名称", "批次名称", "代次", "日期", "接种细胞数", "盘数", "收获细胞数",
                  "本次PD", "累计PD", "状态", "建议下次传代", "备注"], records)
    _write_sheet(workbook, "冻存管",
                 ["ID", "细胞名称", "冻存日期", "细胞数", "位置", "PD", "状态", "备注"], vials)
    _write_sheet(workbook, "细胞组", ["ID", "组名", "备注"], groups)
    _write_sheet(workbook, "组成员", ["组ID", "组名", "细胞名称", "排序"], members)

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def batch_records_csv(conn: sqlite3.Connection, batch_id: int) -> tuple[bytes, str]:
    batch = conn.execute(
        "SELECT cb.batch_name, cl.name AS cell_line FROM culture_batches cb "
        "JOIN cell_lines cl ON cl.id = cb.cell_line_id WHERE cb.id = ?",
        (batch_id,),
    ).fetchone()
    if batch is None:
        raise NotFound(f"培养批次不存在（id={batch_id}）")

    rows = fetch_all(
        conn,
        "SELECT passage, date, inoculum_cells, inoculum_dishes, harvested_cells, pd, cumulative_pd, "
        "status, next_passage_date, notes FROM culture_records WHERE batch_id = ? ORDER BY date ASC, id ASC",
        (batch_id,),
    )
    headers = ["passage", "date", "inoculum_cells", "inoculum_dishes", "harvested_cells",
               "pd", "cumulative_pd", "status", "next_passage_date", "notes"]
    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, fieldnames=headers, lineterminator="\n")
    writer.writeheader()
    for row in rows:
        writer.writerow({key: ("" if row.get(key) is None else row.get(key)) for key in headers})
    # BOM 让 Excel 正确识别 UTF-8 中文
    content = ("\ufeff" + buffer.getvalue()).encode("utf-8")
    filename = f"{batch['batch_name']}_records_{date.today():%Y%m%d}.csv"
    return content, filename
