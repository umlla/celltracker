"""API 端到端测试（Flask test client + 临时数据库）。"""

from __future__ import annotations

import io

import pytest
from openpyxl import Workbook, load_workbook

from backend import db
from backend.services import maintenance

from .conftest import add_record, create_batch, create_line


# --------------------------------------------------------------------------
# 基础与细胞株
# --------------------------------------------------------------------------
def test_meta_and_dashboard_empty(client):
    meta = client.get("/api/meta").get_json()
    assert meta["app_version"].startswith("4.")
    assert meta["counts"]["cell_lines"] == 0
    dashboard = client.get("/api/dashboard").get_json()
    assert dashboard["stats"]["active_batches"] == 0
    assert dashboard["today"]


def test_cell_line_crud_and_conflict(client):
    line = create_line(client, "HeLa", interval=3)
    assert line["name"] == "HeLa"
    assert line["default_passage_interval"] == 3

    duplicate = client.post("/api/cell-lines", json={"name": "HeLa"})
    assert duplicate.status_code == 409

    updated = client.patch(f"/api/cell-lines/{line['id']}", json={"name": "HeLa-S3"})
    assert updated.status_code == 200
    assert updated.get_json()["name"] == "HeLa-S3"

    listing = client.get("/api/cell-lines?search=HeLa").get_json()
    assert [item["name"] for item in listing] == ["HeLa-S3"]

    deleted = client.delete(f"/api/cell-lines/{line['id']}")
    assert deleted.status_code == 200
    assert client.get(f"/api/cell-lines/{line['id']}").status_code == 404


def test_create_line_requires_name(client):
    response = client.post("/api/cell-lines", json={"type": "永生化"})
    assert response.status_code == 400
    assert response.get_json()["error"]["code"] == "validation_error"


# --------------------------------------------------------------------------
# 批次与传代记录
# --------------------------------------------------------------------------
def test_batch_and_record_pd_flow(client):
    line = create_line(client)
    batch = create_batch(client, line["id"])
    assert batch["batch_name"] == f"{line['name']}-260201"

    first = add_record(client, batch["id"])
    assert first["pd"] == pytest.approx(1.0)
    assert first["cumulative_pd"] == pytest.approx(2.0)  # 初始 PD 1 + 1
    assert first["inoculum_cells"] == pytest.approx(2.0)
    assert first["next_passage_date"] == "2026-02-04"
    assert first["is_active"] == 1

    second = add_record(client, batch["id"], date="2026-02-04", harvested_cells=8.0,
                        next_passage_days=0)
    assert second["pd"] == pytest.approx(2.0)
    assert second["cumulative_pd"] == pytest.approx(4.0)

    detail = client.get(f"/api/batches/{batch['id']}").get_json()
    assert len(detail["records"]) == 2
    assert detail["records"][0]["per_dish_cells"] == pytest.approx(1.0)
    active = [record for record in detail["records"] if record["is_active"] == 1]
    assert len(active) == 1
    assert active[0]["id"] == second["id"]
    assert detail["summary"]["cumulative_pd"] == pytest.approx(4.0)

    # 编辑第 1 条记录后整批重算：pd1=2, cum1=3；pd2=2, cum2=5
    updated = client.patch(f"/api/records/{first['id']}", json={
        "harvested_cells": 8.0, "per_dish_cells": 1.0, "next_passage_days": 0,
    })
    assert updated.status_code == 200
    records = client.get(f"/api/batches/{batch['id']}/records").get_json()
    assert [record["cumulative_pd"] for record in records] == pytest.approx([3.0, 5.0])

    # 删除第 2 条后，活跃记录回到第 1 条
    assert client.delete(f"/api/records/{second['id']}").status_code == 200
    records = client.get(f"/api/batches/{batch['id']}/records").get_json()
    assert len(records) == 1
    assert records[0]["is_active"] == 1

    # 终止批次后不能再加记录
    terminated = client.post(f"/api/batches/{batch['id']}/terminate", json={"reason": "污染"})
    assert terminated.status_code == 200
    assert terminated.get_json()["is_active"] == 0
    assert client.post("/api/records", json={
        "batch_id": batch["id"], "harvested_cells": 1,
    }).status_code == 409


def test_record_with_negative_pd_keeps_cumulative(client):
    line = create_line(client)
    batch = create_batch(client, line["id"])
    add_record(client, batch["id"], harvested_cells=4.0)  # pd=1, cum=2
    negative = add_record(client, batch["id"], date="2026-02-05", harvested_cells=1.0)
    assert negative["pd"] == pytest.approx(-1.0)  # log2(1 / 上次接种 2)
    assert negative["cumulative_pd"] == pytest.approx(2.0)


def test_bulk_create_reports_per_row_errors(client):
    line = create_line(client)
    batch_a = create_batch(client, line["id"])
    batch_b = create_batch(client, line["id"], revival_date="2026-03-01")

    response = client.post("/api/records/bulk", json={"entries": [
        {"batch_id": batch_a["id"], "date": "2026-02-02", "harvested_cells": 4.0,
         "inoculum_dishes": 1, "per_dish_cells": 2.0},
        {"batch_id": batch_b["id"], "date": "2026-02-02", "harvested_cells": None},
    ]})
    body = response.get_json()
    assert response.status_code == 200
    assert body["success"] == 1
    assert body["failed"] == 1
    assert body["results"][1]["ok"] is False
    assert len(client.get(f"/api/batches/{batch_b['id']}/records").get_json()) == 0


def test_bulk_update_recalculates(client):
    line = create_line(client)
    batch = create_batch(client, line["id"])
    first = add_record(client, batch["id"])
    second = add_record(client, batch["id"], date="2026-02-06", harvested_cells=6.0)

    response = client.post("/api/records/bulk-update", json={"entries": [
        {"id": first["id"], "harvested_cells": 4.0, "per_dish_cells": 1.0},
        {"id": second["id"], "harvested_cells": 16.0},
    ]})
    assert response.status_code == 200
    assert response.get_json()["success"] == 2
    records = client.get(f"/api/batches/{batch['id']}/records").get_json()
    # pd1 = 1, cum=2；pd2 = log2(16/2)=3，cum=5
    assert [record["cumulative_pd"] for record in records] == pytest.approx([2.0, 5.0])


def test_record_can_create_freeze_vials(client):
    line = create_line(client, "冻存联动细胞", type_="永生化")
    batch = create_batch(client, line["id"])
    record = add_record(
        client,
        batch["id"],
        harvested_cells=4.0,
        freeze_vials={"enabled": True, "cell_count_per_vial": 5.0, "count": 2, "include_pd": True},
    )
    assert record["vials_count"] == 2
    assert len(record["vials_created"]) == 2

    vials = client.get(f"/api/vials?cell_line_id={line['id']}").get_json()
    assert len(vials) == 2
    for vial in vials:
        assert vial["freeze_date"] == "2026-02-02"  # 与传代日期一致
        assert vial["cell_count"] == pytest.approx(5.0)
        assert vial["pd"] == pytest.approx(2.0)  # 初始 PD 1 + 本次 1
        assert vial["location"] == ""
        assert vial["status"] == "在库"
        assert vial["batch_id"] == batch["id"]
        assert "P1" in vial["notes"]


def test_immortalized_can_skip_pd_but_primary_cannot(client):
    immortal = create_line(client, "永生化细胞", type_="永生化")
    immortal_batch = create_batch(client, immortal["id"])
    record = add_record(
        client,
        immortal_batch["id"],
        freeze_vials={"enabled": True, "cell_count_per_vial": 1.0, "count": 1, "include_pd": False},
    )
    assert record["vials_count"] == 1
    vial = client.get(f"/api/vials?cell_line_id={immortal['id']}").get_json()[0]
    assert vial["pd"] is None

    primary = create_line(client, "原代细胞", type_="原代")
    primary_batch = create_batch(client, primary["id"])
    primary_record = add_record(
        client,
        primary_batch["id"],
        freeze_vials={"enabled": True, "cell_count_per_vial": 2.0, "count": 1, "include_pd": False},
    )
    assert primary_record["vials_count"] == 1
    primary_vial = client.get(f"/api/vials?cell_line_id={primary['id']}").get_json()[0]
    assert primary_vial["pd"] == pytest.approx(2.0)


def test_freeze_vials_validation_and_bulk(client):
    line = create_line(client, "批量冻存细胞")
    batch = create_batch(client, line["id"])

    missing = client.post("/api/records", json={
        "batch_id": batch["id"], "harvested_cells": 4.0,
        "freeze_vials": {"enabled": True, "count": 1},
    })
    assert missing.status_code == 400
    assert "每管冻存细胞数" in missing.get_json()["error"]["message"]

    zero_count = client.post("/api/records", json={
        "batch_id": batch["id"], "harvested_cells": 4.0,
        "freeze_vials": {"enabled": True, "cell_count_per_vial": 1.0, "count": 0},
    })
    assert zero_count.status_code == 400

    bulk = client.post("/api/records/bulk", json={"entries": [
        {"batch_id": batch["id"], "date": "2026-02-03", "harvested_cells": 8.0,
         "freeze_vials": {"enabled": True, "cell_count_per_vial": 3.0, "count": 2}},
        {"batch_id": batch["id"], "date": "2026-02-05", "harvested_cells": 8.0},
    ]})
    assert bulk.status_code == 200
    body = bulk.get_json()
    assert body["success"] == 2
    assert body["vials_created"] == 2
    assert body["results"][0]["vials_created"] == 2
    assert body["results"][1]["vials_created"] == 0
    assert len(client.get(f"/api/vials?cell_line_id={line['id']}").get_json()) == 2


def test_batch_csv_export(client):
    line = create_line(client)
    batch = create_batch(client, line["id"])
    add_record(client, batch["id"])
    response = client.get(f"/api/batches/{batch['id']}/records.csv")
    assert response.status_code == 200
    assert response.data.startswith("\ufeff".encode("utf-8"))
    text = response.data.decode("utf-8-sig")
    assert "passage,date,inoculum_cells" in text


def test_delete_batch_and_line_cascade(client):
    line = create_line(client, "级联细胞")
    batch = create_batch(client, line["id"])
    add_record(client, batch["id"])
    assert client.delete(f"/api/batches/{batch['id']}").status_code == 200
    assert client.get(f"/api/batches/{batch['id']}").status_code == 404

    batch2 = create_batch(client, line["id"], revival_date="2026-04-01")
    add_record(client, batch2["id"])
    assert client.delete(f"/api/cell-lines/{line['id']}").status_code == 200
    assert client.get("/api/cell-lines").get_json() == []


# --------------------------------------------------------------------------
# 冻存管 / 冻存盒
# --------------------------------------------------------------------------
def test_vials_and_boxes(client):
    line = create_line(client, "Vial细胞")
    created = client.post("/api/vials", json={
        "cell_line_id": line["id"],
        "freeze_date": "2026-01-05",
        "cell_count": 5.1e5,
        "pd": 3.2,
        "location": "罐A/架1/盒1/A1, 罐A/架1/盒1/A2",
        "count": 1,
    })
    assert created.status_code == 201
    vials = created.get_json()
    assert len(vials) == 2
    assert {vial["location"].split("/")[-1] for vial in vials} == {"A1", "A2"}

    boxes = client.get("/api/vials/boxes").get_json()
    assert boxes[0]["box"] == "罐A/架1/盒1"
    assert boxes[0]["occupied"] == 2

    detail = client.get("/api/vials/boxes/detail?box=罐A/架1/盒1").get_json()
    assert set(detail["holes"].keys()) == {"A1", "A2"}
    assert detail["holes"]["A1"]["cell_name"] == "Vial细胞"

    updated = client.patch(f"/api/vials/{vials[0]['id']}", json={"status": "已丢弃"})
    assert updated.get_json()["status"] == "已丢弃"
    assert client.delete(f"/api/vials/{vials[1]['id']}").status_code == 200
    assert len(client.get("/api/vials").get_json()) == 1


def test_vials_unplaced_filter(client):
    line = create_line(client, "待定位细胞")
    client.post("/api/vials", json={"cell_line_id": line["id"], "cell_count": 1.0, "location": ""})
    client.post("/api/vials", json={"cell_line_id": line["id"], "cell_count": 2.0,
                                    "location": "罐Z/架1/盒1/A1"})
    all_vials = client.get(f"/api/vials?cell_line_id={line['id']}").get_json()
    assert len(all_vials) == 2
    unplaced = client.get(f"/api/vials?cell_line_id={line['id']}&unplaced=1").get_json()
    assert len(unplaced) == 1
    assert unplaced[0]["location"] == ""


def test_batch_binds_vial_and_marks_revived(client):
    line = create_line(client, "复苏细胞")
    vial = client.post("/api/vials", json={
        "cell_line_id": line["id"], "freeze_date": "2026-01-05",
        "cell_count": 1e6, "location": "罐B/架2/盒3/C4",
    }).get_json()[0]

    batch = create_batch(client, line["id"], vial_id=vial["id"])
    refreshed = client.get(f"/api/vials?cell_line_id={line['id']}").get_json()[0]
    assert refreshed["status"] == "已复苏"
    assert refreshed["batch_id"] == batch["id"]

    reuse = client.post("/api/batches", json={
        "cell_line_id": line["id"], "vial_id": vial["id"], "revival_date": "2026-02-02",
    })
    assert reuse.status_code == 409


def test_vial_excel_import_append_and_replace(client):
    create_line(client, "HeLa")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "冻存盒A"
    sheet["A1"] = "HeLa\n5X10^5\nPD=10.26\n2022.4.15"
    sheet["B1"] = "新细胞\n2X10^6\n2023.5.1"
    buffer = io.BytesIO()
    workbook.save(buffer)
    payload = buffer.getvalue()

    def upload(mode: str):
        return client.post(
            "/api/vials/import",
            data={"file": (io.BytesIO(payload), "vials.xlsx"), "mode": mode},
            content_type="multipart/form-data",
        )

    response = upload("append")
    assert response.status_code == 200, response.get_json()
    result = response.get_json()
    assert result["inserted"] == 2
    assert result["lines_created"] == ["新细胞"]
    boxes = client.get("/api/vials/boxes").get_json()
    assert [box["box"] for box in boxes] == ["冻存盒A"]

    again = upload("append").get_json()
    assert again["inserted"] == 0
    assert again["skipped"] == 2

    replaced = upload("replace").get_json()
    assert replaced["inserted"] == 2
    assert len(client.get("/api/vials").get_json()) == 2


# --------------------------------------------------------------------------
# 细胞组
# --------------------------------------------------------------------------
def test_group_members_order_status_and_curves(client):
    line_a = create_line(client, "A细胞")
    line_b = create_line(client, "B细胞")
    group = client.post("/api/groups", json={"group_name": "实验组", "notes": "测试"}).get_json()
    assert group["member_count"] == 0

    members = client.put(f"/api/groups/{group['id']}/members",
                         json={"members": [line_b["id"], line_a["id"]]})
    assert members.status_code == 200
    assert [member["name"] for member in members.get_json()["members"]] == ["B细胞", "A细胞"]

    batch_a = create_batch(client, line_a["id"])
    batch_b = create_batch(client, line_b["id"])
    add_record(client, batch_a["id"], date="2026-02-02", harvested_cells=8.0)
    add_record(client, batch_b["id"], date="2026-02-02", harvested_cells=4.0)
    add_record(client, batch_b["id"], date="2026-02-04", harvested_cells=8.0)

    status = client.get(f"/api/groups/{group['id']}/status").get_json()
    assert len(status) == 2
    assert all(row["remaining_days"] is not None for row in status)

    curves = client.get(f"/api/groups/{group['id']}/curves").get_json()
    assert len(curves["series"]) == 2
    assert curves["dates"] == ["2026-02-02", "2026-02-04"]

    normalized = client.get(
        f"/api/groups/{group['id']}/curves?normalize=true&start_date=2026-02-02"
    ).get_json()
    assert len(normalized["series"]) == 2
    for series in normalized["series"]:
        first_point = series["points"][0]
        assert first_point["value"] == pytest.approx(1.0)

    rates = client.get(f"/api/groups/{group['id']}/rates").get_json()
    # 只有 B 细胞有两条记录（间隔 2 天，PD=2），A 只有一条无法计算速率
    assert {row["cell_line"] for row in rates} == {"B细胞"}
    for row in rates:
        assert row["avg_pd_per_day"] == pytest.approx(1.0)  # PD 2 / 2 天


def test_group_bulk_batch_create_and_terminate(client):
    line_a = create_line(client, "组批A")
    line_b = create_line(client, "组批B")
    group = client.post("/api/groups", json={
        "group_name": "批量组", "cell_line_ids": [line_a["id"], line_b["id"]],
    }).get_json()

    created = client.post(f"/api/groups/{group['id']}/batches", json={
        "revival_date": "2026-05-01", "initial_pd": 0.5, "initial_cell_count": 1.0,
    })
    assert created.status_code == 201
    assert len(created.get_json()["created"]) == 2

    terminated = client.post(f"/api/groups/{group['id']}/terminate", json={"reason": "实验结束"})
    assert terminated.status_code == 200
    assert terminated.get_json()["count"] == 2
    assert client.get(f"/api/groups/{group['id']}/status").get_json() == []


def test_group_records_endpoint(client):
    line = create_line(client, "组记录细胞")
    group = client.post("/api/groups", json={
        "group_name": "记录组", "cell_line_ids": [line["id"]],
    }).get_json()
    batch = create_batch(client, line["id"])
    add_record(client, batch["id"])

    rows = client.get(f"/api/groups/{group['id']}/records").get_json()
    assert len(rows) == 1
    assert rows[0]["cell_line"] == "组记录细胞"
    assert rows[0]["batch_name"] == batch["batch_name"]
    assert rows[0]["per_dish_cells"] == pytest.approx(1.0)


def test_group_export_xlsx(client):
    line = create_line(client, "导出细胞")
    group = client.post("/api/groups", json={
        "group_name": "导出组", "cell_line_ids": [line["id"]],
    }).get_json()
    batch = create_batch(client, line["id"])
    add_record(client, batch["id"])

    response = client.get(f"/api/groups/{group['id']}/export?include_all=true")
    assert response.status_code == 200
    workbook = load_workbook(io.BytesIO(response.data))
    assert workbook.sheetnames == ["细胞株信息", "培养批次", "传代记录", "冻存管"]
    assert workbook["传代记录"].max_row == 2
    assert workbook["细胞株信息"]["B1"].value == "细胞名称"


def test_full_export(client):
    line = create_line(client, "全量细胞")
    response = client.get("/api/export/full")
    assert response.status_code == 200
    workbook = load_workbook(io.BytesIO(response.data))
    assert "细胞组" in workbook.sheetnames
    assert workbook["细胞株信息"]["B2"].value == "全量细胞"


# --------------------------------------------------------------------------
# 维护
# --------------------------------------------------------------------------
def test_audit_and_recalculate(client):
    line = create_line(client)
    batch = create_batch(client, line["id"])
    add_record(client, batch["id"])
    audit = client.get("/api/maintenance/audit").get_json()
    assert audit["issues"] == []
    assert audit["pd_audit"]["mismatch_count"] == 0

    recalculated = client.post("/api/maintenance/recalculate")
    assert recalculated.status_code == 200
    assert recalculated.get_json()["remaining_issues"] == []


def test_backup_and_restore(client, config):
    line = create_line(client, "备份细胞")
    batch = create_batch(client, line["id"])
    add_record(client, batch["id"])

    response = client.get("/api/maintenance/backup")
    assert response.status_code == 200
    assert response.data[:16] == b"SQLite format 3\x00"

    backup_path = config.backup_dir / "restore-source.db"
    backup_path.write_bytes(response.data)

    # 破坏数据后恢复
    client.delete(f"/api/cell-lines/{line['id']}")
    assert client.get("/api/cell-lines").get_json() == []

    with open(backup_path, "rb") as handle:
        restored = client.post(
            "/api/maintenance/restore",
            data={"file": (handle, "restore-source.db")},
            content_type="multipart/form-data",
        )
    assert restored.status_code == 200, restored.get_json()
    lines = client.get("/api/cell-lines").get_json()
    assert [item["name"] for item in lines] == ["备份细胞"]
    assert client.get("/api/maintenance/audit").get_json()["pd_audit"]["mismatch_count"] == 0


def test_import_legacy_endpoint_validates_and_restores(client):
    create_line(client, "导入源细胞")
    backup = client.get("/api/maintenance/backup").data
    assert backup[:16] == b"SQLite format 3\x00"

    response = client.post(
        "/api/maintenance/import-legacy",
        data={"file": (io.BytesIO(backup), "legacy.db")},
        content_type="multipart/form-data",
    )
    assert response.status_code == 200, response.get_json()
    report = response.get_json()
    assert report["counts_equal"] is True
    assert report["checksums_equal"] is True
    assert report["status"] == "ok"
    assert [line["name"] for line in client.get("/api/cell-lines").get_json()] == ["导入源细胞"]

    # 非法文件被拒绝，原数据保持不变
    invalid = client.post(
        "/api/maintenance/import-legacy",
        data={"file": (io.BytesIO(b"not a sqlite file"), "broken.db")},
        content_type="multipart/form-data",
    )
    assert invalid.status_code == 400
    assert [line["name"] for line in client.get("/api/cell-lines").get_json()] == ["导入源细胞"]


def test_migration_report_endpoint(client, config):
    assert client.get("/api/maintenance/migration-report").get_json() is None
    report_path = config.db_path.parent / "migration-report.json"
    report_path.write_text('{"status": "ok"}', encoding="utf-8")
    assert client.get("/api/maintenance/migration-report").get_json()["status"] == "ok"


def test_backup_listing(client, config):
    maintenance.create_backup(config)
    backups = client.get("/api/maintenance/backups").get_json()
    assert len(backups) == 1
    assert backups[0]["name"].endswith(".db")
