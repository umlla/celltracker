import streamlit as st
import sqlite3
import pandas as pd
import plotly.graph_objects as go
from datetime import datetime, timedelta, date
import math
import openpyxl
import re
# ==================== 数据库初始化 ====================
DB_NAME = "cell_data.db"

def init_db():
    conn = sqlite3.connect(DB_NAME)
    c = conn.cursor()

    c.execute('''
        CREATE TABLE IF NOT EXISTS cell_lines (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL UNIQUE,
            revival_date TEXT,
            type TEXT,
            default_passage_interval INTEGER DEFAULT 0,
            notes TEXT
        )
    ''')

    c.execute('''
    CREATE TABLE IF NOT EXISTS culture_batches (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        cell_line_id INTEGER,
        batch_name TEXT NOT NULL,
        revival_date TEXT,
        source_freeze_date TEXT,
        initial_pd REAL DEFAULT 0,
        initial_cell_count REAL,
        is_active INTEGER DEFAULT 1,
        end_date TEXT,
        end_reason TEXT,
        notes TEXT,
        FOREIGN KEY(cell_line_id) REFERENCES cell_lines(id)
    )
    ''')

    c.execute('''
        CREATE TABLE IF NOT EXISTS culture_records (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            batch_id INTEGER,
            passage INTEGER,
            date TEXT,
            inoculum_cells REAL,
            inoculum_dishes INTEGER,
            harvested_cells REAL,
            pd REAL,
            cumulative_pd REAL,
            status TEXT,
            next_passage_date TEXT,
            is_active INTEGER DEFAULT 1,
            notes TEXT,
            FOREIGN KEY(batch_id) REFERENCES culture_batches(id)
        )
    ''')

    c.execute('''
        CREATE TABLE IF NOT EXISTS frozen_vials (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            cell_line_id INTEGER,
            batch_id INTEGER,
            freeze_date TEXT,
            cell_count REAL,
            location TEXT DEFAULT '',
            pd REAL,
            status TEXT DEFAULT '在库',
            notes TEXT,
            FOREIGN KEY(cell_line_id) REFERENCES cell_lines(id),
            FOREIGN KEY(batch_id) REFERENCES culture_batches(id)
        )
    ''')

    c.execute('''
        CREATE TABLE IF NOT EXISTS cell_groups (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            group_name TEXT NOT NULL UNIQUE,
            notes TEXT
        )
    ''')

    c.execute('''
        CREATE TABLE IF NOT EXISTS cell_group_members (
            group_id INTEGER,
            cell_line_id INTEGER,
            sort_order INTEGER DEFAULT 0,
            FOREIGN KEY(group_id) REFERENCES cell_groups(id),
            FOREIGN KEY(cell_line_id) REFERENCES cell_lines(id),
            PRIMARY KEY(group_id, cell_line_id)
        )
    ''')

    conn.commit()
    conn.close()

init_db()
def upgrade_db():
    conn = sqlite3.connect(DB_NAME)
    c = conn.cursor()
    # ... 其他已有升级 ...
    # 新增排序字段
    try:
        c.execute("ALTER TABLE cell_group_members ADD COLUMN sort_order INTEGER DEFAULT 0")
    except sqlite3.OperationalError:
        pass
    conn.commit()
    conn.close()
upgrade_db()

# ==================== 辅助函数 ====================


def get_connection():
    return sqlite3.connect(DB_NAME)

def calculate_pd(harvested, inoculum):
    if inoculum and inoculum > 0 and harvested and harvested > 0:
        return round(math.log2(harvested / inoculum), 2)
    return 0

def get_batch_active_record(batch_id):
    conn = get_connection()
    df = pd.read_sql_query(
        "SELECT * FROM culture_records WHERE batch_id = ? AND is_active = 1 ORDER BY date DESC LIMIT 1",
        conn, params=(batch_id,))
    conn.close()
    return df.iloc[0] if not df.empty else None

def get_batch_all_records(batch_id):
    conn = get_connection()
    df = pd.read_sql_query(
        "SELECT * FROM culture_records WHERE batch_id = ? ORDER BY date ASC",
        conn, params=(batch_id,))
    conn.close()
    return df

def get_cell_line_frozen_vials(cell_line_id):
    conn = get_connection()
    df = pd.read_sql_query(
        "SELECT * FROM frozen_vials WHERE cell_line_id = ? ORDER BY freeze_date DESC",
        conn, params=(cell_line_id,))
    conn.close()
    return df

def get_cell_line_info(cell_line_id):
    conn = get_connection()
    df = pd.read_sql_query("SELECT * FROM cell_lines WHERE id = ?", conn, params=(cell_line_id,))
    conn.close()
    return df.iloc[0] if not df.empty else None

def get_all_cell_lines():
    conn = get_connection()
    df = pd.read_sql("SELECT id, name, type FROM cell_lines ORDER BY name", conn)
    conn.close()
    return df

def get_groups():
    conn = get_connection()
    df = pd.read_sql("SELECT * FROM cell_groups ORDER BY group_name", conn)
    conn.close()
    return df

def get_group_members(group_id):
    conn = get_connection()
    # 检查列是否存在，若不存在则添加
    cols = [row[1] for row in conn.execute("PRAGMA table_info(cell_group_members)").fetchall()]
    if 'sort_order' not in cols:
        conn.execute("ALTER TABLE cell_group_members ADD COLUMN sort_order INTEGER DEFAULT 0")
        conn.commit()
    df = pd.read_sql_query('''
        SELECT cl.id, cl.name, cl.type, gm.sort_order
        FROM cell_group_members gm
        JOIN cell_lines cl ON gm.cell_line_id = cl.id
        WHERE gm.group_id = ?
        ORDER BY gm.sort_order, cl.name
    ''', conn, params=(group_id,))
    conn.close()
    return df
def get_batch_last_record(batch_id):
    conn = get_connection()
    df = pd.read_sql_query(
        "SELECT * FROM culture_records WHERE batch_id = ? ORDER BY date DESC LIMIT 1",
        conn, params=(batch_id,))
    conn.close()
    return df.iloc[0] if not df.empty else None
def recalculate_cumulative_pd(batch_id):
    conn = get_connection()
    batch = pd.read_sql_query(
        "SELECT initial_pd, initial_cell_count FROM culture_batches WHERE id = ?",
        conn, params=(batch_id,))
    if batch.empty:
        conn.close()
        return
    initial_pd = batch.iloc[0]['initial_pd'] or 0.0
    initial_cell_count = batch.iloc[0]['initial_cell_count']

    records = pd.read_sql_query(
        "SELECT * FROM culture_records WHERE batch_id = ? ORDER BY date ASC, id ASC",
        conn, params=(batch_id,))

    cumulative = initial_pd
    prev_inoculum = initial_cell_count   # None 作为首次分母
    for _, row in records.iterrows():
        harvested = row['harvested_cells']
        if prev_inoculum is not None and prev_inoculum > 0 and harvested > 0:
            pd_val = round(math.log2(harvested / prev_inoculum), 2)
        else:
            pd_val = 0
        cumulative += max(0, pd_val)
        conn.execute("UPDATE culture_records SET pd = ?, cumulative_pd = ? WHERE id = ?",
                     (pd_val, round(cumulative, 2), row['id']))
        prev_inoculum = row['inoculum_cells']
    conn.commit()
    conn.close()
def parse_number(s):
    """解析各种数字格式，支持科学计数简写 (如 5.1X105, 1.2*10^6)"""
    s = s.strip().replace(',', '').replace(' ', '')
    m = re.match(r'(\d+(?:\.\d+)?)\s*[xX×\*]\s*10\s*\^?\s*(\d+)', s)
    if m:
        base = float(m.group(1))
        exp = int(m.group(2))
        return base * (10 ** exp)
    try:
        return float(s)
    except:
        pass
    return None

def parse_date(date_str):
    """更强大的日期解析，优先用 pandas，支持 2022.4.15 等无前导零格式"""
    s = date_str.strip()
    if not s:
        return None
    try:
        dt = pd.to_datetime(s, dayfirst=False, yearfirst=False)
        if pd.notna(dt):
            return dt.strftime('%Y-%m-%d')
    except:
        pass
    for fmt in ['%Y-%m-%d', '%Y/%m/%d', '%Y.%m.%d', '%d/%m/%Y', '%m/%d/%Y',
                '%Y-%m', '%Y/%m', '%Y.%m']:
        try:
            dt = datetime.strptime(s, fmt)
            return dt.strftime('%Y-%m-%d')
        except:
            pass
    try:
        int(s)
        if len(s) == 4:
            return s + '-01-01'
    except:
        pass
    return None

def parse_cell_text(text):
    """解析单元格文本，支持多行、单行混合信息"""
    if not text or not str(text).strip():
        return None
    raw = str(text).strip()
    lines = [l.strip() for l in raw.split('\n') if l.strip()]
    if len(lines) == 1 and ' ' in lines[0]:
        alt = re.split(r'\s{2,}|\t|,|;', lines[0])
        if len(alt) > 1:
            lines = alt
        else:
            lines = lines[0].split()
            lines = [l for l in lines if l]
    if not lines:
        return None
    cell_name = lines[0]
    numbers = []
    dates = []
    others = []
    for token in lines[1:]:
        token = token.strip()
        if not token:
            continue
        pd_match = re.search(r'[pP][dD]\s*[=：:]\s*(\d+(?:\.\d+)?)', token)
        if pd_match:
            numbers.append(float(pd_match.group(1)))
            continue
        d = parse_date(token)
        if d:
            dates.append(d)
            continue
        val = parse_number(token)
        if val is not None:
            numbers.append(val)
        else:
            others.append(token)
    cell_count = numbers[0] if numbers else None
    pd_val = numbers[1] if len(numbers) > 1 else None
    freeze_date = dates[0] if dates else None
    revived_date = dates[1] if len(dates) > 1 else None
    status = '已复苏' if revived_date else '在库'
    notes = ' '.join(others)
    return {
        'cell_name': cell_name,
        'cell_count': cell_count,
        'pd': pd_val,
        'freeze_date': freeze_date,
        'revived_date': revived_date,
        'status': status,
        'notes': notes
    }
# ==================== 页面配置 ====================
st.set_page_config(page_title="细胞培养追踪系统 v3.0", layout="wide")
st.title("🧬 细胞培养追踪系统 v3.0")

menu = st.sidebar.radio("功能菜单", [
    "🧫 在养细胞",
    "➕ 添加细胞株",
    "📂 管理细胞组",
    "🧬 新建传代记录",
    "✏️ 编辑传代记录（按组）",
    "❄️ 添加冻存管",
    "🧊 冻存盒视图", 
    "🧪 新建培养批次"
])
# ==================== 一次性修复累计PD（仅管理员） ====================
with st.sidebar.expander("🔧 管理工具"):
    if st.button("修复所有批次累计PD（忽略负值）"):
        conn = get_connection()
        batch_ids = pd.read_sql_query("SELECT id FROM culture_batches", conn)['id'].tolist()
        fixed = 0
        for bid in batch_ids:
            recalculate_cumulative_pd(bid)  # 确保该函数已修改为忽略负PD
            fixed += 1
        conn.close()
        st.success(f"已修复 {fixed} 个批次的累计PD。")
        st.rerun()
# ==================== 1. 在养细胞 ====================
if menu == "🧫 在养细胞":
    st.header("在养细胞总览")

    # 计算在养细胞株总数（有不重复活跃批次）
    conn = get_connection()
    active_line_count = pd.read_sql_query('''
        SELECT COUNT(DISTINCT cl.id) as cnt
        FROM culture_batches cb
        JOIN cell_lines cl ON cb.cell_line_id = cl.id
        WHERE cb.is_active = 1
    ''', conn).iloc[0, 0]
    conn.close()
    st.metric("在养细胞株总数", active_line_count)

    view_mode = st.radio("查看模式", ["按细胞株查看", "按细胞组查看"])

    if view_mode == "按细胞株查看":
        all_lines = get_all_cell_lines()
        if all_lines.empty:
            st.info("暂无细胞株，请先添加。")
        else:
            # ---------- 使用 session_state 避免删除/改名后崩溃 ----------
            if 'selected_line_id' not in st.session_state:
                st.session_state.selected_line_id = None

            all_names = all_lines["name"].tolist()
            if st.session_state.selected_line_id is not None:
                mask = all_lines["id"] == st.session_state.selected_line_id
                if mask.any():
                    default_name = all_lines.loc[mask, "name"].values[0]
                    default_index = all_names.index(default_name)
                else:
                    default_index = 0
                    st.session_state.selected_line_id = None
            else:
                default_index = 0

            selected_name = st.selectbox("选择细胞株（可输入搜索）", all_names, index=default_index)
            selected_id = int(all_lines[all_lines["name"] == selected_name]["id"].iloc[0])
            st.session_state.selected_line_id = selected_id

            line_info = get_cell_line_info(selected_id)

            # 显示基本信息
            st.subheader(f"细胞株: {line_info['name']}")
            st.caption(f"类型: {line_info['type']}  |  默认传代间隔: {line_info['default_passage_interval']} 天")

            # 可选显示已终止批次
            show_ended = st.checkbox("显示已终止的批次", key=f"show_ended_{selected_id}")

            conn = get_connection()
            if show_ended:
                batches_df = pd.read_sql_query(
                    "SELECT * FROM culture_batches WHERE cell_line_id = ? ORDER BY is_active DESC, revival_date DESC",
                    conn, params=(selected_id,))
            else:
                batches_df = pd.read_sql_query(
                    "SELECT * FROM culture_batches WHERE cell_line_id = ? AND is_active = 1",
                    conn, params=(selected_id,))
            conn.close()

            st.write("---")
            st.subheader("🧪 培养批次")
            if batches_df.empty:
                st.warning("没有符合条件的批次。")
            else:
                for _, batch in batches_df.iterrows():
                    batch_id = batch['id']
                    is_alive = batch['is_active'] == 1

                    # 获取用于展示的记录
                    if is_alive:
                        active_record = get_batch_active_record(batch_id)
                        display_record = active_record
                    else:
                        active_record = None
                        display_record = get_batch_last_record(batch_id)

                    # 构建标题
                    if display_record is not None:
                        title = f"📌 {batch['batch_name']} (P{display_record['passage']})"
                    else:
                        title = f"⚪ {batch['batch_name']}"

                    if not is_alive:
                        title += f"  🔴 已终止（{batch['end_reason']}，{batch['end_date']}）"

                    with st.expander(title):
                        if display_record is not None:
                            col1, col2, col3, col4 = st.columns(4)
                            col1.metric("当前盘数", display_record["inoculum_dishes"])
                            col2.metric("细胞总数", f'{display_record["inoculum_cells"]:,.0f}')
                            col3.metric("传代次数", f'P{display_record["passage"]}')
                            col4.metric("累计PD", display_record["cumulative_pd"])
                            st.write(f"**状态**: {display_record['status'] or '—'}")
                            st.write(f"**传代时间**: {display_record['date']}")
                            if display_record["next_passage_date"]:
                                st.write(f"**建议下次传代时间**: {display_record['next_passage_date']}")

                           

                            with st.expander("📜 查看所有传代记录与生长曲线"):
                                all_records = get_batch_all_records(batch_id)
                                if not all_records.empty:
                                    st.caption(f"总传代次数：{len(all_records)}")
                                    st.dataframe(
                                        all_records[["passage","date","inoculum_cells","harvested_cells","pd","cumulative_pd"]]
                                        .style.format({
                                            "passage": "{:.0f}",
                                            "inoculum_cells": "{:.2f}",
                                            "harvested_cells": "{:.2f}",
                                            "pd": "{:.2f}",
                                            "cumulative_pd": "{:.2f}"
                                        }),
                                        width='stretch'
                                    )
                                    fig = go.Figure()
                                    fig.add_trace(go.Scatter(x=all_records["date"], y=all_records["cumulative_pd"],
                                                             mode="lines+markers", name="累计PD", yaxis="y1",
                                                             line=dict(width=2)))
                                    fig.add_trace(go.Scatter(x=all_records["date"], y=all_records["pd"],
                                                             mode="lines+markers", name="单次PD", yaxis="y1",
                                                             line=dict(dash="dot", width=1.5, color="#17becf")))
                                    fig.add_trace(go.Scatter(x=all_records["date"], y=all_records["harvested_cells"],
                                                             mode="lines+markers", name="收获细胞数", yaxis="y2",
                                                             line=dict(dash="dot", color="#ff7f0e")))
                                    fig.update_layout(
                                        xaxis=dict(title="日期"),
                                        yaxis=dict(title="PD值", side="left"),
                                        yaxis2=dict(title="细胞总数 (×10⁵ cells)", overlaying="y", side="right", type="log"),
                                        hovermode="x unified",
                                        legend=dict(x=0.01, y=0.99),
                                        margin=dict(l=60, r=60)
                                    )
                                    st.plotly_chart(fig, width='stretch', key=f"growth_curve_{batch_id}")
                                    csv = all_records.to_csv(index=False).encode('utf-8')
                                    st.download_button("📥 导出记录为CSV", csv, f"{batch['batch_name']}_records.csv", "text/csv",
                                                       key=f"download_csv_{batch_id}")
                                else:
                                    st.info("暂无传代记录。")

                        # ----- 终止此批次培养（仅活跃批次） -----
                        if is_alive:
                            st.write("---")
                            with st.expander("⏹️ 终止此批次培养"):
                                end_reason = st.selectbox(
                                    "终止原因",
                                    ["污染", "全部冻存", "丢弃", "实验结束", "其他"],
                                    key=f"end_reason_{batch_id}"
                                )
                                if st.button("确认终止", key=f"end_btn_{batch_id}"):
                                    conn = get_connection()
                                    today_str = date.today().isoformat()
                                    conn.execute(
                                        "UPDATE culture_batches SET is_active = 0, end_date = ?, end_reason = ? WHERE id = ?",
                                        (today_str, end_reason, batch_id)
                                    )
                                    conn.execute(
                                        "UPDATE culture_records SET is_active = 0 WHERE batch_id = ? AND is_active = 1",
                                        (batch_id,)
                                    )
                                    conn.commit()
                                    conn.close()
                                    st.success(f"批次「{batch['batch_name']}」已终止（原因：{end_reason}）")
                                    st.rerun()
                            st.write("---")
                            with st.expander("🗑 删除此批次（危险操作）"):
                                st.error("此操作将永久删除该批次及其所有传代记录，且不可恢复！")
                                del_batch_confirm = st.checkbox(
                                    "我确认要删除此批次及所有传代记录",
                                    key=f"del_batch_confirm_{batch_id}"
                                )
                                if del_batch_confirm:
                                    if st.button("删除批次", key=f"btn_del_batch_{batch_id}"):
                                        conn = get_connection()
                                        conn.execute("DELETE FROM culture_records WHERE batch_id = ?", (batch_id,))
                                        conn.execute("UPDATE frozen_vials SET batch_id = NULL WHERE batch_id = ?", (batch_id,))
                                        conn.execute("DELETE FROM culture_batches WHERE id = ?", (batch_id,))
                                        conn.commit()
                                        conn.close()
                                        st.success(f"批次「{batch['batch_name']}」及其传代记录已删除。")
                                        st.rerun()

            # 冻存状态
            st.subheader("❄️ 冻存状态")
            vials = get_cell_line_frozen_vials(selected_id)
            if not vials.empty:
                total_vials = len(vials)
                total_cells = vials["cell_count"].sum()
                st.write(f"总冻存管数: {total_vials}  |  总冻存细胞数: {total_cells:,.0f}")
                vials_display = vials[["freeze_date", "cell_count", "location", "pd", "status"]].copy()
                vials_display["pd"] = vials_display["pd"].apply(
                    lambda x: "" if pd.isna(x) else f"{x:.2f}"
                )
                st.dataframe(vials_display, width='stretch',
                            column_config={"status": "状态"})
            else:
                st.info("暂无冻存记录。")

            # 编辑细胞株信息
            st.write("---")
            with st.expander("✏️ 编辑细胞株信息"):
                with st.form("edit_line_in_view"):
                    new_name = st.text_input("细胞株名称", value=line_info["name"])
                    new_interval = st.number_input("默认传代间隔（天）", min_value=0,
                                                   value=int(line_info["default_passage_interval"]))
                    new_revival = st.date_input("复苏日期",
                                                value=datetime.strptime(line_info["revival_date"], "%Y-%m-%d").date()
                                                if line_info["revival_date"] else datetime.today())
                    new_notes = st.text_area("备注", value=line_info["notes"] if line_info["notes"] else "")
                    if st.form_submit_button("保存修改"):
                        if new_name != line_info["name"]:
                            conn = get_connection()
                            conflict = conn.execute("SELECT id FROM cell_lines WHERE name = ? AND id != ?",
                                                    (new_name, selected_id)).fetchone()
                            conn.close()
                            if conflict:
                                st.error("该名称已被其他细胞株使用，请更换。")
                                st.stop()
                        conn = get_connection()
                        conn.execute("UPDATE cell_lines SET name = ?, default_passage_interval = ?, revival_date = ?, notes = ? WHERE id = ?",
                                     (new_name, new_interval, new_revival.isoformat(), new_notes, selected_id))
                        conn.commit()
                        conn.close()
                        st.success("信息已更新")
                        st.session_state.selected_line_id = None   # ← 清除选中缓存
                        st.rerun()

            # 删除细胞株
            st.write("---")
            st.warning("⚠️ 危险操作：删除细胞株将同时删除其所有批次、传代记录、冻存管及分组关系。")
            del_cell_confirm = st.checkbox("我确认要删除此细胞株及其所有数据", key=f"del_cell_{selected_id}")
            if del_cell_confirm:
                if st.button("删除此细胞株", key=f"btn_del_cell_{selected_id}"):
                    conn = get_connection()
                    conn.execute("DELETE FROM frozen_vials WHERE cell_line_id = ?", (selected_id,))
                    conn.execute("DELETE FROM culture_records WHERE batch_id IN (SELECT id FROM culture_batches WHERE cell_line_id = ?)", (selected_id,))
                    conn.execute("DELETE FROM culture_batches WHERE cell_line_id = ?", (selected_id,))
                    conn.execute("DELETE FROM cell_group_members WHERE cell_line_id = ?", (selected_id,))
                    conn.execute("DELETE FROM cell_lines WHERE id = ?", (selected_id,))
                    conn.commit()
                    conn.close()
                    st.success("细胞株已删除。")
                    st.session_state.selected_line_id = None   # ← 清除选中缓存
                    st.rerun()

    else:  # 按细胞组查看
        groups_df = get_groups()
        if groups_df.empty:
            st.info("暂无细胞组，请先创建。")
        else:
            group_name = st.selectbox("选择细胞组", groups_df["group_name"].tolist())
            group_id = int(groups_df[groups_df["group_name"] == group_name]["id"].iloc[0])
            members_df = get_group_members(group_id)
            if members_df.empty:
                st.warning("该组暂无成员。")
            else:
                member_ids = members_df["id"].tolist()

                # ========== 组内批次批量操作（新增） ==========
                with st.expander("⚡ 组内批次批量操作"):
                    op = st.radio("选择操作", ["终止该组所有批次", "批量新建培养批次"], horizontal=True)
                    if op == "终止该组所有批次":
                        end_reason = st.selectbox("统一终止原因", ["污染", "全部冻存", "丢弃", "实验结束", "其他"], key="group_end_reason")
                        if st.button("执行终止", key="group_terminate"):
                            conn = get_connection()
                            today_str = date.today().isoformat()
                            placeholders = ','.join(['?']*len(member_ids))
                            # 查找该组所有活跃批次
                            batch_df = pd.read_sql_query(
                                f"SELECT id FROM culture_batches WHERE is_active=1 AND cell_line_id IN ({placeholders})",
                                conn, params=member_ids
                            )
                            if not batch_df.empty:
                                batch_ids = batch_df['id'].tolist()
                                for bid in batch_ids:
                                    conn.execute("UPDATE culture_batches SET is_active=0, end_date=?, end_reason=? WHERE id=?",
                                                 (today_str, end_reason, bid))
                                    conn.execute("UPDATE culture_records SET is_active=0 WHERE batch_id=? AND is_active=1", (bid,))
                                conn.commit()
                                st.success(f"已终止 {len(batch_ids)} 个批次。")
                            else:
                                st.info("该组无活跃批次。")
                            conn.close()
                            st.rerun()
                    else:   # 批量新建培养批次
                        with st.form("group_new_batches"):
                            st.subheader("为组内所有细胞株新建培养批次")
                            revival_date = st.date_input("复苏日期", value=datetime.today())
                            initial_pd = st.number_input("初始PD值（统一）", min_value=0.0, value=0.0, format="%.2f")
                            initial_cell_count = st.number_input("复苏接种细胞数（统一，可选）", min_value=0.0, value=0.0, format="%.2f")
                            source_freeze_date = st.date_input("冻存来源日期（可选）", value=datetime.today())
                            notes = st.text_area("批次备注（统一）")
                            if st.form_submit_button("创建批次"):
                                conn = get_connection()
                                created = 0
                                for _, member in members_df.iterrows():
                                    line_id = member["id"]
                                    line_name = member["name"]
                                    batch_name = f"{line_name}-{revival_date.strftime('%y%m%d')}"
                                    conn.execute('''INSERT INTO culture_batches
                                        (cell_line_id, batch_name, revival_date, source_freeze_date,
                                         initial_pd, initial_cell_count, is_active, notes)
                                        VALUES (?,?,?,?,?,?,1,?)''',
                                        (line_id, batch_name, revival_date.isoformat(), source_freeze_date.isoformat(),
                                         initial_pd, initial_cell_count if initial_cell_count > 0 else None, notes))
                                    created += 1
                                conn.commit()
                                conn.close()
                                st.success(f"已为 {created} 个细胞株创建新批次。")
                                st.rerun()

                # ========== 原有表格查询（状态一览） ==========
                conn = get_connection()
                placeholders = ','.join(['?'] * len(member_ids))
                query = f'''
                    SELECT cb.id as batch_id, cb.batch_name, cl.name as cell_name, cl.type,
                           cr.passage, cr.date as last_passage, cr.inoculum_cells, cr.inoculum_dishes,
                           cr.cumulative_pd, cr.status, cr.next_passage_date
                    FROM culture_batches cb
                    JOIN cell_lines cl ON cb.cell_line_id = cl.id
                    LEFT JOIN culture_records cr ON cr.batch_id = cb.id
                        AND cr.is_active = 1
                        AND cr.date = (SELECT MAX(date) FROM culture_records WHERE batch_id = cb.id AND is_active = 1)
                    WHERE cb.is_active = 1 AND cl.id IN ({placeholders})
                    ORDER BY cl.name, cb.batch_name
                '''
                df = pd.read_sql_query(query, conn, params=member_ids)
                conn.close()

                if df.empty:
                    st.info("组内无正在培养的批次。")
                else:
                    today = date.today()
                    df['delta'] = df['next_passage_date'].apply(
                        lambda x: (datetime.strptime(x, "%Y-%m-%d").date() - today).days if pd.notnull(x) else None)
                    def color_delta(val):
                        if val is None or pd.isna(val):
                            return ''
                        elif val <= 0:
                            return 'background-color: #ffcccc'
                        elif val <= 2:
                            return 'background-color: #fff3cd'
                        else:
                            return 'background-color: #d4edda'
                    styled_df = df.style.map(color_delta, subset=['delta'])
                    st.dataframe(styled_df, column_config={
                        "batch_name": "批次", "cell_name": "细胞株", "type": "类型",
                        "passage": "代次", "last_passage": "上次传代",
                        "inoculum_dishes": "盘数", "inoculum_cells": "细胞总数",
                        "cumulative_pd": "累计PD", "status": "状态",
                        "next_passage_date": "建议下次传代", "delta": "剩余天数"
                    }, width='stretch')
                    # ---------- 已终止批次选择（可多选加入生长曲线） ----------
                    conn = get_connection()
                    ended_batches = pd.read_sql_query(f'''
                        SELECT cb.id as batch_id, cb.batch_name, cl.name as cell_name
                        FROM culture_batches cb
                        JOIN cell_lines cl ON cb.cell_line_id = cl.id
                        WHERE cb.is_active = 0 AND cl.id IN ({placeholders})
                        ORDER BY cl.name, cb.batch_name
                    ''', conn, params=member_ids)
                    conn.close()
                    
                    selected_ended_ids = []
                    if not ended_batches.empty:
                        ended_options = [f"{r['cell_name']} - {r['batch_name']}" for _, r in ended_batches.iterrows()]
                        ended_ids = ended_batches['batch_id'].tolist()
                        selected_ended = st.multiselect(
                            "📦 选择已终止批次加入生长曲线",
                            ended_options,
                            key=f"ended_select_{group_id}"
                        )
                        selected_ended_ids = [ended_ids[i] for i, opt in enumerate(ended_options) if opt in selected_ended]
                    # -------------------------------------------------------------
                    # ========== 增殖速率均值对比 ==========
                    st.subheader("📊 增殖速率对比（最近三次 PD/天 均值）")
                    batch_ids = df['batch_id'].tolist()
                    batch_names = df['batch_name'].tolist()
                    conn = get_connection()
                    avg_rates = []
                    for bid, bname in zip(batch_ids, batch_names):
                        records = pd.read_sql_query(
                            "SELECT date, pd FROM culture_records WHERE batch_id = ? ORDER BY date",
                            conn, params=(bid,)
                        )
                        if len(records) >= 2:
                            records['date_dt'] = pd.to_datetime(records['date'])
                            records['interval_days'] = records['date_dt'].diff().dt.days
                            records['pd_per_day'] = records['pd'] / records['interval_days']
                            rates = records['pd_per_day'].dropna().tolist()
                            if rates:
                                last_rates = rates[-3:]
                                avg_rate = sum(last_rates) / len(last_rates)
                                avg_rates.append({
                                    'batch': bname,
                                    'avg_pd_per_day': avg_rate
                                })
                    conn.close()

                    if avg_rates:
                        avg_df = pd.DataFrame(avg_rates)
                        fig_avg = go.Figure()
                        fig_avg.add_trace(go.Bar(
                            x=avg_df['batch'],
                            y=avg_df['avg_pd_per_day'],
                            marker_color='#2ca02c'
                        ))
                        fig_avg.update_layout(
                            xaxis_title="批次",
                            yaxis_title="平均 PD/天",
                            hovermode="x"
                        )
                        st.plotly_chart(fig_avg, width='stretch', key="group_avg_rate")
                    else:
                        st.info("组内尚无足够传代记录用于计算速率均值。")

                    # ========== 组内所有批次的生长曲线（一张图） ==========
                    st.subheader("📈 组内生长曲线（累计PD vs 日期）")
                    
                    # 标准化开关
                    use_normalize = st.checkbox("标准化到指定代数起点（起点=1）", value=False, key=f"norm_{group_id}")
                    start_passage = 1
                    if use_normalize:
                        # 收集所有批次（活跃 + 已终止）
                        all_records = []
                        conn = get_connection()
                        for _, batch_row in df.iterrows():
                            bid = batch_row['batch_id']
                            bname = batch_row['batch_name']
                            records = pd.read_sql_query(
                                "SELECT date, cumulative_pd FROM culture_records WHERE batch_id = ? ORDER BY date",
                                conn, params=(bid,)
                            )
                            if not records.empty:
                                all_records.append((bname, records, False))
                        for eid in selected_ended_ids:
                            e_row = ended_batches[ended_batches['batch_id'] == eid].iloc[0]
                            e_name = f"{e_row['cell_name']} - {e_row['batch_name']} (已终止)"
                            e_records = pd.read_sql_query(
                                "SELECT date, cumulative_pd FROM culture_records WHERE batch_id = ? ORDER BY date",
                                conn, params=(eid,)
                            )
                            if not e_records.empty:
                                all_records.append((e_name, e_records, True))
                        conn.close()

                        # 获取所有日期用于选择器
                        all_dates = set()
                        for _, recs, _ in all_records:
                            for d in recs['date']:
                                all_dates.add(d)
                        sorted_dates = sorted(all_dates)

                        if sorted_dates:
                            start_date = st.selectbox("选择起始日期（仅显示当天有记录的细胞）", sorted_dates, key=f"start_date_{group_id}")
                        else:
                            start_date = None

                        fig_group = go.Figure()
                        any_curve = False
                        for bname, records, is_ended in all_records:
                            # 严格匹配起始日期
                            exact_match = records[records['date'] == start_date]
                            if exact_match.empty:
                                continue
                            # 从该日期开始的所有记录
                            sub = records[records['date'] >= start_date].copy()
                            if len(sub) < 1:
                                continue
                            base_pd = exact_match.iloc[0]['cumulative_pd']
                            if base_pd == 0:
                                continue
                            sub['normalized_pd'] = sub['cumulative_pd'] / base_pd
                            dash_style = 'dash' if is_ended else None
                            label = f"{bname} (起始 {start_date})"
                            fig_group.add_trace(go.Scatter(
                                x=sub["date"],
                                y=sub["normalized_pd"],
                                mode="lines+markers",
                                line=dict(dash=dash_style) if dash_style else {},
                                name=label
                            ))
                            any_curve = True

                        if any_curve:
                            fig_group.update_layout(
                                xaxis_title="日期",
                                yaxis_title="标准化累计PD (起点=1)",
                                hovermode="x unified",
                                legend=dict(title="批次")
                            )
                            st.plotly_chart(fig_group, width='stretch', key=f"group_growth_curve_norm_{group_id}")
                        else:
                            st.info("所选日期没有任何细胞有传代记录。")
                    else:
                        # ===== 默认曲线：原始累计PD（活跃批次 + 选中的已终止批次） =====
                        fig_group = go.Figure()
                        conn = get_connection()
                        # 活跃批次
                        for _, batch_row in df.iterrows():
                            bid = batch_row['batch_id']
                            bname = batch_row['batch_name']
                            records = pd.read_sql_query(
                                "SELECT date, cumulative_pd FROM culture_records WHERE batch_id = ? ORDER BY date",
                                conn, params=(bid,)
                            )
                            if not records.empty:
                                fig_group.add_trace(go.Scatter(
                                    x=records["date"],
                                    y=records["cumulative_pd"],
                                    mode="lines+markers",
                                    name=bname
                                ))
                        # 选中的已终止批次（虚线）
                        for eid in selected_ended_ids:
                            e_row = ended_batches[ended_batches['batch_id'] == eid].iloc[0]
                            e_name = f"{e_row['cell_name']} - {e_row['batch_name']} (已终止)"
                            e_records = pd.read_sql_query(
                                "SELECT date, cumulative_pd FROM culture_records WHERE batch_id = ? ORDER BY date",
                                conn, params=(eid,)
                            )
                            if not e_records.empty:
                                fig_group.add_trace(go.Scatter(
                                    x=e_records["date"],
                                    y=e_records["cumulative_pd"],
                                    mode="lines+markers",
                                    line=dict(dash='dash'),
                                    name=e_name
                                ))
                        conn.close()

                        if len(fig_group.data) > 0:
                            fig_group.update_layout(
                                xaxis_title="日期",
                                yaxis_title="累计PD",
                                hovermode="x unified",
                                legend=dict(title="批次")
                            )
                            st.plotly_chart(fig_group, width='stretch', key=f"group_growth_curve_{group_id}")
                        else:
                            st.info("组内尚无传代记录。")

                # ========== 批量导出组数据（可读版，保持原样） ==========
                st.subheader("📥 批量导出组数据")
                include_all = st.checkbox("包含已终止的批次和所有传代记录", value=True, key=f"export_all_{group_id}")
                if st.button("导出组内全部数据为 Excel", key=f"export_btn_{group_id}"):
                    conn = get_connection()
                    lines_query = f'''
                        SELECT id, name AS 细胞名称, type AS 类型, revival_date AS 复苏日期,
                               default_passage_interval AS 默认传代间隔, notes AS 备注
                        FROM cell_lines
                        WHERE id IN ({placeholders})
                    '''
                    df_lines = pd.read_sql_query(lines_query, conn, params=member_ids)

                    if include_all:
                        batches_query = f'''
                            SELECT cb.id, cl.name AS 细胞名称, cb.batch_name AS 批次名称,
                                   cb.revival_date AS 复苏日期, cb.source_freeze_date AS 冻存来源日期,
                                   cb.initial_pd AS 初始PD, cb.initial_cell_count AS 初始细胞数,
                                   cb.is_active AS 是否活跃, cb.end_date AS 终止日期, cb.end_reason AS 终止原因,
                                   cb.notes AS 备注
                            FROM culture_batches cb
                            JOIN cell_lines cl ON cb.cell_line_id = cl.id
                            WHERE cl.id IN ({placeholders})
                        '''
                    else:
                        batches_query = f'''
                            SELECT cb.id, cl.name AS 细胞名称, cb.batch_name AS 批次名称,
                                   cb.revival_date AS 复苏日期, cb.source_freeze_date AS 冻存来源日期,
                                   cb.initial_pd AS 初始PD, cb.initial_cell_count AS 初始细胞数,
                                   cb.is_active AS 是否活跃, cb.end_date AS 终止日期, cb.end_reason AS 终止原因,
                                   cb.notes AS 备注
                            FROM culture_batches cb
                            JOIN cell_lines cl ON cb.cell_line_id = cl.id
                            WHERE cl.id IN ({placeholders}) AND cb.is_active = 1
                        '''
                    df_batches = pd.read_sql_query(batches_query, conn, params=member_ids)

                    if df_batches.empty:
                        st.warning("没有批次数据可导出。")
                        conn.close()
                    else:
                        batch_ids = df_batches["id"].tolist()
                        batch_placeholders = ','.join(['?'] * len(batch_ids))
                        if include_all:
                            records_query = f'''
                                SELECT cr.id, cl.name AS 细胞名称, cb.batch_name AS 批次名称,
                                       cr.passage AS 代次, cr.date AS 日期,
                                       cr.inoculum_cells AS 接种细胞数, cr.inoculum_dishes AS 盘数,
                                       cr.harvested_cells AS 收获细胞数, cr.pd AS 本次PD,
                                       cr.cumulative_pd AS 累计PD, cr.status AS 状态,
                                       cr.next_passage_date AS 建议下次传代, cr.notes AS 备注
                                FROM culture_records cr
                                JOIN culture_batches cb ON cr.batch_id = cb.id
                                JOIN cell_lines cl ON cb.cell_line_id = cl.id
                                WHERE cr.batch_id IN ({batch_placeholders})
                            '''
                        else:
                            records_query = f'''
                                SELECT cr.id, cl.name AS 细胞名称, cb.batch_name AS 批次名称,
                                       cr.passage AS 代次, cr.date AS 日期,
                                       cr.inoculum_cells AS 接种细胞数, cr.inoculum_dishes AS 盘数,
                                       cr.harvested_cells AS 收获细胞数, cr.pd AS 本次PD,
                                       cr.cumulative_pd AS 累计PD, cr.status AS 状态,
                                       cr.next_passage_date AS 建议下次传代, cr.notes AS 备注
                                FROM culture_records cr
                                JOIN culture_batches cb ON cr.batch_id = cb.id
                                JOIN cell_lines cl ON cb.cell_line_id = cl.id
                                WHERE cr.batch_id IN ({batch_placeholders}) AND cr.is_active = 1
                            '''
                        df_records = pd.read_sql_query(records_query, conn, params=batch_ids)

                        vials_query = f'''
                            SELECT fv.id, cl.name AS 细胞名称, fv.freeze_date AS 冻存日期,
                                   fv.cell_count AS 细胞数, fv.location AS 位置, fv.pd AS PD,
                                   fv.status AS 状态, fv.notes AS 备注
                            FROM frozen_vials fv
                            JOIN cell_lines cl ON fv.cell_line_id = cl.id
                            WHERE cl.id IN ({placeholders})
                        '''
                        df_vials = pd.read_sql_query(vials_query, conn, params=member_ids)
                        conn.close()

                        from io import BytesIO
                        output = BytesIO()
                        with pd.ExcelWriter(output, engine='openpyxl') as writer:
                            df_lines.to_excel(writer, sheet_name='细胞株信息', index=False)
                            df_batches.to_excel(writer, sheet_name='培养批次', index=False)
                            df_records.to_excel(writer, sheet_name='传代记录', index=False)
                            df_vials.to_excel(writer, sheet_name='冻存管', index=False)

                        file_name = f"细胞组数据_{group_name}.xlsx"
                        st.download_button(
                            label="⬇️ 下载 Excel 文件",
                            data=output.getvalue(),
                            file_name=file_name,
                            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                        )
                        st.success("文件生成成功，请点击上方按钮下载。")

# ==================== 2. 添加细胞株 ====================
elif menu == "➕ 添加细胞株":
    st.header("添加新细胞株")
    with st.form("add_line"):
        name = st.text_input("细胞株名称 *")
        revival_date = st.date_input("首次复苏日期", value=datetime.today())
        line_type = st.selectbox("类型", ["永生化", "原代"])
        default_interval = st.number_input("默认传代间隔（天，0表示不设默认）", min_value=0, value=0, step=1)
        notes = st.text_area("备注")
        if st.form_submit_button("添加"):
            if not name:
                st.error("名称不能为空")
            else:
                conn = get_connection()
                try:
                    conn.execute("INSERT INTO cell_lines (name, revival_date, type, default_passage_interval, notes) VALUES (?,?,?,?,?)",
                                 (name, revival_date.isoformat(), line_type, default_interval, notes))
                    conn.commit()
                    st.success(f"已添加细胞株「{name}」")
                except sqlite3.IntegrityError:
                    st.error("该细胞株名称已存在。")
                finally:
                    conn.close()

# ==================== 3. 管理细胞组 ====================
elif menu == "📂 管理细胞组":
    st.header("管理细胞组")
    tab1, tab2 = st.tabs(["创建/删除组", "管理组成员"])

    with tab1:
        st.subheader("创建新细胞组")
        with st.form("add_group"):
            group_name = st.text_input("组名")
            notes = st.text_area("备注")
            if st.form_submit_button("创建"):
                if group_name:
                    conn = get_connection()
                    try:
                        conn.execute("INSERT INTO cell_groups (group_name, notes) VALUES (?,?)", (group_name, notes))
                        conn.commit()
                        st.success(f"已创建组「{group_name}」")
                    except sqlite3.IntegrityError:
                        st.error("组名已存在。")
                    finally:
                        conn.close()
                else:
                    st.error("组名不能为空。")

        groups_df = get_groups()
        if not groups_df.empty:
            st.subheader("现有组（点击删除）")
            for _, group in groups_df.iterrows():
                col1, col2 = st.columns([4,1])
                col1.write(f"**{group['group_name']}**")
                if col2.button("删除", key=f"del_{group['id']}"):
                    conn = get_connection()
                    conn.execute("DELETE FROM cell_group_members WHERE group_id = ?", (group['id'],))
                    conn.execute("DELETE FROM cell_groups WHERE id = ?", (group['id'],))
                    conn.commit()
                    conn.close()
                    st.rerun()

        with tab2:
            st.subheader("修改组成员与排序")
            groups_df = get_groups()
            if groups_df.empty:
                st.info("请先创建细胞组。")
            else:
                group_name = st.selectbox("选择细胞组", groups_df["group_name"].tolist())
                group_id = int(groups_df[groups_df["group_name"] == group_name]["id"].iloc[0])
                members_df = get_group_members(group_id)

                # ---------- 成员选择（不变） ----------
                all_lines = get_all_cell_lines()
                all_names = all_lines["name"].tolist() if not all_lines.empty else []
                current_members = members_df["name"].tolist()
                new_members = st.multiselect("组成员（勾选添加，取消勾选移除）",
                                            options=all_names, default=current_members)
                if st.button("保存成员更改"):
                    to_add = list(set(new_members) - set(current_members))
                    to_remove = list(set(current_members) - set(new_members))
                    conn = get_connection()
                    # 添加新成员，sort_order 取当前最大值+1
                    max_order = conn.execute("SELECT MAX(sort_order) FROM cell_group_members WHERE group_id = ?",
                                            (group_id,)).fetchone()[0] or 0
                    for name in to_add:
                        line_id = int(all_lines[all_lines["name"] == name]["id"].iloc[0])
                        max_order += 1
                        conn.execute("INSERT OR IGNORE INTO cell_group_members (group_id, cell_line_id, sort_order) VALUES (?,?,?)",
                                    (group_id, line_id, max_order))
                    for name in to_remove:
                        line_id = int(all_lines[all_lines["name"] == name]["id"].iloc[0])
                        conn.execute("DELETE FROM cell_group_members WHERE group_id = ? AND cell_line_id = ?",
                                    (group_id, line_id))
                    conn.commit()
                    conn.close()
                    st.success("成员已更新。")
                    st.rerun()

                         # ---------- 排序调整 ----------
                        # ---------- 排序调整 ----------
                if not members_df.empty:
                    if 'sort_order' not in members_df.columns:
                        members_df['sort_order'] = 0
                    st.markdown("---")
                    st.subheader("调整排序（修改数字后点击保存）")

                    # 准备用于编辑的表格
                    order_df = members_df[['id', 'name', 'sort_order']].copy()
                    order_df['sort_order'] = order_df['sort_order'].astype(int)

                    edited_order = st.data_editor(
                        order_df,
                        column_config={
                            "id": None,  # 隐藏ID列
                            "name": st.column_config.TextColumn("细胞株", disabled=True),
                            "sort_order": st.column_config.NumberColumn("排序号", min_value=0, step=1)
                        },
                        width='stretch',
                        hide_index=True,
                        num_rows="fixed"
                    )

                    if st.button("保存排序"):
                        conn = get_connection()
                        for _, row in edited_order.iterrows():
                            conn.execute(
                                "UPDATE cell_group_members SET sort_order = ? WHERE group_id = ? AND cell_line_id = ?",
                                (int(row["sort_order"]), group_id, int(row["id"]))
                            )
                        conn.commit()
                        conn.close()
                        st.success("排序已更新")
                        st.rerun()


# ==================== 4. 新建传代记录 ====================
elif menu == "🧬 新建传代记录":
    st.header("新建传代记录")
    mode = st.radio("操作模式", ["按组录入（表格式）", "单批次录入", "按组编辑（表格式）"], horizontal=True, index=0)

    # ========================= 按组录入（表格式） =========================
    if mode == "按组录入（表格式）":
        groups_df = get_groups()
        if groups_df.empty:
            st.warning("暂无细胞组，请先创建或切换到单批次录入。")
        else:
            group_name = st.selectbox("选择细胞组", groups_df["group_name"].tolist())
            group_id = int(groups_df[groups_df["group_name"] == group_name]["id"].iloc[0])
            members_df = get_group_members(group_id)
            if members_df.empty:
                st.warning("该组暂无成员。")
            else:
                # 构建表格数据：每个细胞株的活跃批次，预填默认值
                table_data = []
                today_str = datetime.today().strftime("%Y-%m-%d")
                for _, member in members_df.iterrows():
                    line_id = member["id"]
                    line_name = member["name"]
                    line_info = get_cell_line_info(line_id)
                    default_interval = line_info["default_passage_interval"]

                    conn = get_connection()
                    batches = pd.read_sql_query(
                        "SELECT id, batch_name FROM culture_batches WHERE cell_line_id = ? AND is_active = 1",
                        conn, params=(line_id,))
                    conn.close()

                    if batches.empty:
                        continue
                    for _, batch in batches.iterrows():
                        batch_id = batch["id"]
                        batch_name = batch["batch_name"]
                        # 获取当前最大代次
                        conn = get_connection()
                        max_p = pd.read_sql_query(
                            "SELECT MAX(passage) FROM culture_records WHERE batch_id = ?",
                            conn, params=(batch_id,)).iloc[0, 0]
                        conn.close()
                        next_p = int(max_p) + 1 if max_p is not None else 1
                        table_data.append({
                            "细胞株": line_name,
                            "批次": batch_name,
                            "批次ID": batch_id,
                            "是否传代": True,          # 新增：默认勾选
                            "传代次数": next_p,
                            "传代日期": today_str,
                            "收获细胞总数": 0.0,
                            "种下盘数": 1,
                            "每盘细胞数": 0.0,
                            "细胞状态": "良好",
                            "建议间隔天数": default_interval,
                            "备注": ""
                        })

                if not table_data:
                    st.info("该组内所有细胞株均无活跃批次，请先创建批次。")
                else:
                    st.subheader(f"批量传代 - {group_name}")
                    df = pd.DataFrame(table_data)

                    column_config = {
                        "细胞株": st.column_config.TextColumn("细胞株", disabled=True, width="small"),
                        "批次": st.column_config.TextColumn("批次", disabled=True, width="small"),
                        "批次ID": None,
                        "是否传代": st.column_config.CheckboxColumn("传代", default=True, width="small"),
                        "传代次数": st.column_config.NumberColumn("传代次数", min_value=1, step=1, width="small"),
                        "传代日期": st.column_config.TextColumn("传代日期", width="small"),
                        "收获细胞总数": st.column_config.NumberColumn("收获总数", min_value=0.0, format="%.2f", width="small"),
                        "种下盘数": st.column_config.NumberColumn("盘数", min_value=1, step=1, width="small"),
                        "每盘细胞数": st.column_config.NumberColumn("每盘细胞数", min_value=0.0, format="%.2f", width="small"),
                        "细胞状态": st.column_config.SelectboxColumn("状态", options=["良好", "一般", "较差", "污染"], width="small"),
                        "建议间隔天数": st.column_config.NumberColumn("间隔天数", min_value=0, step=1, width="small"),
                        "备注": st.column_config.TextColumn("备注", width="small")
                    }

                    edited_df = st.data_editor(
                        df,
                        column_config=column_config,
                        width='stretch',
                        hide_index=True,
                        num_rows="fixed",
                        column_order=["细胞株", "批次", "是否传代", "传代次数", "传代日期", "收获细胞总数",
                                      "种下盘数", "每盘细胞数", "细胞状态", "建议间隔天数", "备注"]
                    )

                    if st.button("批量提交", type="primary"):
                        # 只提交勾选了“是否传代”的行
                        to_submit = edited_df[edited_df["是否传代"] == True]
                        if to_submit.empty:
                            st.warning("没有需要提交的记录，请勾选“传代”列。")
                        else:
                            errors = []
                            success_count = 0
                            for idx, row in to_submit.iterrows():
                                try:
                                    batch_id = int(df.at[idx, "批次ID"])
                                    passage = int(row["传代次数"])
                                    date_str = str(row["传代日期"])[:10]
                                    try:
                                        date_obj = datetime.strptime(date_str, "%Y-%m-%d").date()
                                    except:
                                        errors.append(f"第{idx+1}行日期格式错误，请使用YYYY-MM-DD")
                                        continue
                                    harvested_cells = float(row["收获细胞总数"])
                                    inoculum_dishes = int(row["种下盘数"])
                                    per_dish_cells = float(row["每盘细胞数"])
                                    inoculum_cells = inoculum_dishes * per_dish_cells
                                    status = row["细胞状态"]
                                    next_days = int(row["建议间隔天数"])
                                    notes = row["备注"]

                                    # 计算 PD
                                    prev_active = get_batch_active_record(batch_id)
                                    if prev_active is not None:
                                        prev_inoculum = prev_active["inoculum_cells"]
                                        prev_cumulative_pd = prev_active["cumulative_pd"]
                                    else:
                                        conn = get_connection()
                                        batch_info = pd.read_sql_query(
                                            "SELECT initial_pd, initial_cell_count FROM culture_batches WHERE id = ?",
                                            conn, params=(batch_id,))
                                        conn.close()
                                        initial_pd = batch_info.iloc[0]['initial_pd'] if not batch_info.empty else 0.0
                                        initial_cell_count = batch_info.iloc[0]['initial_cell_count'] if not batch_info.empty else None
                                        prev_inoculum = initial_cell_count if initial_cell_count else 0
                                        prev_cumulative_pd = initial_pd

                                    pd_value = calculate_pd(harvested_cells, prev_inoculum)
                                    cumulative_pd = round(prev_cumulative_pd + max(0, pd_value), 2)

                                    conn = get_connection()
                                    conn.execute("UPDATE culture_records SET is_active = 0 WHERE batch_id = ? AND is_active = 1", (batch_id,))
                                    next_date = None
                                    if next_days > 0:
                                        next_date = (date_obj + timedelta(days=next_days)).isoformat()
                                    conn.execute('''INSERT INTO culture_records
                                        (batch_id, passage, date, inoculum_cells, inoculum_dishes, harvested_cells,
                                         pd, cumulative_pd, status, next_passage_date, is_active, notes)
                                        VALUES (?,?,?,?,?,?,?,?,?,?,1,?)''',
                                        (batch_id, passage, date_obj.isoformat(), inoculum_cells, inoculum_dishes,
                                         harvested_cells, pd_value, cumulative_pd, status, next_date, notes))
                                    conn.commit()
                                    conn.close()
                                    success_count += 1
                                except Exception as e:
                                    errors.append(f"第{idx+1}行提交失败：{e}")

                            if errors:
                                for e in errors:
                                    st.error(e)
                            st.success(f"成功提交 {success_count} 条传代记录。")
                            if not errors:
                                st.rerun()
    elif mode == "按组编辑（表格式）":
        groups_df = get_groups()
        if groups_df.empty:
            st.warning("暂无细胞组，请先创建。")
        else:
            group_name = st.selectbox("选择细胞组", groups_df["group_name"].tolist())
            group_id = int(groups_df[groups_df["group_name"] == group_name]["id"].iloc[0])
            members_df = get_group_members(group_id)
            if members_df.empty:
                st.warning("该组暂无成员。")
            else:
                member_ids = members_df["id"].tolist()
                placeholders = ','.join(['?'] * len(member_ids))

                # 查询组内所有活跃批次的**最新活跃记录**
                query = f'''
                    SELECT cr.id as record_id, cr.batch_id, cr.passage, cr.date,
                           cr.harvested_cells, cr.inoculum_cells, cr.inoculum_dishes,
                           cr.status, cr.next_passage_date, cr.notes,
                           cb.batch_name, cl.name as cell_name
                    FROM culture_records cr
                    JOIN culture_batches cb ON cr.batch_id = cb.id
                    JOIN cell_lines cl ON cb.cell_line_id = cl.id
                    WHERE cb.is_active = 1 AND cr.is_active = 1
                      AND cl.id IN ({placeholders})
                    ORDER BY cl.name, cb.batch_name
                '''
                conn = get_connection()
                records_df = pd.read_sql_query(query, conn, params=member_ids)
                conn.close()

                if records_df.empty:
                    st.info("该组内没有可编辑的活跃传代记录。")
                else:
                    st.subheader(f"批量编辑传代记录 - {group_name}")

                    # 构建编辑用表格，保留隐藏 ID
                    table_data = []
                    for _, row in records_df.iterrows():
                        dishes = int(row['inoculum_dishes'])
                        per_dish = round(row['inoculum_cells'] / dishes, 2) if dishes > 0 else 0.0
                        table_data.append({
                            "record_id": int(row['record_id']),
                            "batch_id": int(row['batch_id']),
                            "细胞株": row['cell_name'],
                            "批次": row['batch_name'],
                            "传代次数": int(row['passage']),
                            "传代日期": row['date'],
                            "收获细胞总数": float(row['harvested_cells']),
                            "种下盘数": dishes,
                            "每盘细胞数": per_dish,
                            "细胞状态": row['status'] if row['status'] in ["良好","一般","较差","污染"] else "良好",
                            "下次传代间隔天数": 0,
                            "备注": row['notes'] if row['notes'] else "",
                            "原有下次传代日期": row['next_passage_date']   # 隐藏，确保未修改时保留
                        })

                    df = pd.DataFrame(table_data)

                    column_config = {
                        "record_id": None,
                        "batch_id": None,
                        "原有下次传代日期": None,
                        "细胞株": st.column_config.TextColumn("细胞株", disabled=True, width="small"),
                        "批次": st.column_config.TextColumn("批次", disabled=True, width="small"),
                        "传代次数": st.column_config.NumberColumn("传代次数", min_value=1, step=1, width="small"),
                        "传代日期": st.column_config.TextColumn("传代日期", width="small"),
                        "收获细胞总数": st.column_config.NumberColumn("收获总数", min_value=0.0, format="%.2f", width="small"),
                        "种下盘数": st.column_config.NumberColumn("盘数", min_value=1, step=1, width="small"),
                        "每盘细胞数": st.column_config.NumberColumn("每盘细胞数", min_value=0.0, format="%.2f", width="small"),
                        "细胞状态": st.column_config.SelectboxColumn("状态", options=["良好","一般","较差","污染"], width="small"),
                        "下次传代间隔天数": st.column_config.NumberColumn("间隔天数", min_value=0, step=1, width="small"),
                        "备注": st.column_config.TextColumn("备注", width="small")
                    }

                    edited_df = st.data_editor(
                        df,
                        column_config=column_config,
                        width='stretch',
                        hide_index=True,
                        num_rows="fixed",
                        column_order=["细胞株","批次","传代次数","传代日期","收获细胞总数",
                                      "种下盘数","每盘细胞数","细胞状态","下次传代间隔天数","备注"]
                    )

                    if st.button("批量提交修改", type="primary"):
                        errors = []
                        success_count = 0
                        conn = get_connection()
                        for idx, edited_row in edited_df.iterrows():
                            try:
                                record_id = int(df.at[idx, "record_id"])
                                batch_id = int(df.at[idx, "batch_id"])

                                new_date_str = str(edited_row["传代日期"])[:10]
                                try:
                                    new_date = datetime.strptime(new_date_str, "%Y-%m-%d").date()
                                except:
                                    errors.append(f"第{idx+1}行日期格式错误，请使用YYYY-MM-DD")
                                    continue
                                new_passage = int(edited_row["传代次数"])
                                new_harvested = float(edited_row["收获细胞总数"])
                                new_dishes = int(edited_row["种下盘数"])
                                new_per_dish = float(edited_row["每盘细胞数"])
                                new_inoculum = new_dishes * new_per_dish
                                new_status = edited_row["细胞状态"]
                                next_days = int(edited_row["下次传代间隔天数"])
                                new_notes = edited_row["备注"]

                                # 计算下次传代日期
                                if next_days > 0:
                                    next_date = (new_date + timedelta(days=next_days)).isoformat()
                                else:
                                    next_date = df.at[idx, "原有下次传代日期"]   # 保留原值（可能为 None）

                                # 更新记录
                                conn.execute('''UPDATE culture_records
                                            SET passage=?, date=?, inoculum_cells=?, inoculum_dishes=?,
                                                harvested_cells=?, status=?, next_passage_date=?, notes=?
                                            WHERE id=?''',
                                            (new_passage, new_date.isoformat(), new_inoculum, new_dishes,
                                             new_harvested, new_status, next_date, new_notes, record_id))
                                conn.commit()
                                # 重新计算整批累积 PD
                                recalculate_cumulative_pd(batch_id)
                                success_count += 1
                            except Exception as e:
                                errors.append(f"第{idx+1}行提交失败：{e}")
                        conn.close()

                        if errors:
                            for e in errors:
                                st.error(e)
                        st.success(f"成功修改 {success_count} 条记录。")
                        if not errors:
                            st.rerun()

    # ========================= 单批次录入（保留原逻辑） =========================
    else:
        all_lines = get_all_cell_lines()
        if all_lines.empty:
            st.warning("请先添加细胞株和批次。")
        else:
            line_name = st.selectbox("选择细胞株", all_lines["name"].tolist())
            selected_line_id = int(all_lines[all_lines["name"] == line_name]["id"].iloc[0])
            line_info = get_cell_line_info(selected_line_id)
            default_interval = line_info["default_passage_interval"]

            conn = get_connection()
            batches_df = pd.read_sql_query(
                "SELECT id, batch_name FROM culture_batches WHERE cell_line_id = ? AND is_active = 1",
                conn, params=(selected_line_id,))
            conn.close()

            if batches_df.empty:
                st.error("该细胞株没有活跃批次，请先创建批次。")
            else:
                batch_choice = st.selectbox("选择培养批次", batches_df["batch_name"].tolist())
                selected_batch_id = int(batches_df[batches_df["batch_name"] == batch_choice]["id"].iloc[0])

                # 获取默认代次
                conn = get_connection()
                max_passage = pd.read_sql_query(
                    "SELECT MAX(passage) FROM culture_records WHERE batch_id = ?",
                    conn, params=(selected_batch_id,)).iloc[0, 0]
                conn.close()
                default_passage = int(max_passage) + 1 if max_passage is not None else 1

                with st.form("add_culture"):
                    passage = st.number_input("传代次数 (P)", min_value=1, value=default_passage, step=1)
                    date_input = st.date_input("传代日期", value=datetime.today())
                    harvested_cells = st.number_input("本次收获细胞总数（×10⁵ cells）", min_value=0.0, format="%.2f")
                    inoculum_dishes = st.number_input("本次传代种下盘数", min_value=1, value=1)
                    per_dish_cells = st.number_input("每盘细胞数（×10⁵ cells）", min_value=0.0, format="%.2f")
                    inoculum_cells = inoculum_dishes * per_dish_cells
                    status = st.selectbox("细胞状态", ["良好", "一般", "较差", "污染"])
                    next_days = st.number_input("建议下次传代间隔天数",
                                                min_value=0,
                                                value=default_interval if default_interval > 0 else 0)
                    notes = st.text_area("备注")
                    if st.form_submit_button("提交"):
                        conn = get_connection()
                        batch_info = pd.read_sql_query(
                            "SELECT initial_pd, initial_cell_count FROM culture_batches WHERE id = ?",
                            conn, params=(selected_batch_id,))
                        conn.close()
                        initial_pd = batch_info.iloc[0]['initial_pd'] if not batch_info.empty else 0.0
                        initial_cell_count = batch_info.iloc[0]['initial_cell_count'] if not batch_info.empty else None

                        prev_active = get_batch_active_record(selected_batch_id)
                        if prev_active is not None:
                            prev_inoculum = prev_active["inoculum_cells"]
                            prev_cumulative_pd = prev_active["cumulative_pd"]
                        else:
                            prev_inoculum = initial_cell_count if initial_cell_count else 0
                            prev_cumulative_pd = initial_pd

                        pd_value = calculate_pd(harvested_cells, prev_inoculum)
                        cumulative_pd = round(prev_cumulative_pd + max(0, pd_value), 2)

                        conn = get_connection()
                        conn.execute("UPDATE culture_records SET is_active = 0 WHERE batch_id = ? AND is_active = 1",
                                     (selected_batch_id,))
                        next_date = None
                        if next_days > 0:
                            next_date = (date_input + timedelta(days=next_days)).isoformat()
                        conn.execute('''INSERT INTO culture_records
                            (batch_id, passage, date, inoculum_cells, inoculum_dishes, harvested_cells,
                             pd, cumulative_pd, status, next_passage_date, is_active, notes)
                            VALUES (?,?,?,?,?,?,?,?,?,?,1,?)''',
                            (selected_batch_id, passage, date_input.isoformat(), inoculum_cells, inoculum_dishes,
                             harvested_cells, pd_value, cumulative_pd, status, next_date, notes))
                        conn.commit()
                        conn.close()
                        st.success("传代记录已添加。")
# ==================== 编辑传代记录（按组） ====================
elif menu == "✏️ 编辑传代记录（按组）":
    st.header("按组编辑已有传代记录")
    groups_df = get_groups()
    if groups_df.empty:
        st.warning("请先创建细胞组。")
    else:
        group_name = st.selectbox("选择细胞组", groups_df["group_name"].tolist())
        group_id = int(groups_df[groups_df["group_name"] == group_name]["id"].iloc[0])
        members_df = get_group_members(group_id)
        if members_df.empty:
            st.warning("该组暂无成员。")
        else:
            # 收集组内所有活跃批次的传代记录
            member_ids = members_df["id"].tolist()
            conn = get_connection()
            # 获取活跃批次列表
            batches_df = pd.read_sql_query(
                f"SELECT id, batch_name, cell_line_id FROM culture_batches WHERE is_active = 1 AND cell_line_id IN ({','.join(['?']*len(member_ids))})",
                conn, params=member_ids)
            conn.close()
            if batches_df.empty:
                st.info("组内没有活跃的培养批次。")
            else:
                # 获取每个批次的最新记录（用于展示当前状态）以及所有记录（用于编辑）
                # 我们选择显示每个批次的所有传代记录，按日期排序，让用户编辑任意一条
                all_records_list = []
                for _, batch_row in batches_df.iterrows():
                    bid = batch_row["id"]
                    bname = batch_row["batch_name"]
                    conn = get_connection()
                    records = pd.read_sql_query(
                        "SELECT id, passage, date, inoculum_cells, inoculum_dishes, harvested_cells, status, next_passage_date, notes "
                        "FROM culture_records WHERE batch_id = ? ORDER BY date",
                        conn, params=(bid,))
                    conn.close()
                    records["batch_name"] = bname
                    records["batch_id"] = bid
                    all_records_list.append(records)

                if not all_records_list:
                    st.info("组内暂无传代记录。")
                else:
                    combined_df = pd.concat(all_records_list, ignore_index=True)
                    # 计算每盘细胞数方便编辑
                    combined_df["per_dish_cells"] = combined_df["inoculum_cells"] / combined_df["inoculum_dishes"]
                    combined_df["next_days"] = 0  # 临时列，用于计算间隔天数

                    st.subheader(f"编辑 {group_name} 的传代记录")
                    st.caption("修改任意单元格，点击“批量保存”提交所有修改。")

                    # 转换为可编辑格式
                    edited_df = st.data_editor(
                        combined_df[[
                            "id", "batch_name", "passage", "date", "harvested_cells",
                            "inoculum_dishes", "per_dish_cells", "status", "next_days", "notes"
                        ]],
                        column_config={
                            "id": None,  # 隐藏记录ID
                            "batch_name": st.column_config.TextColumn("批次", disabled=True, width="small"),
                            "passage": st.column_config.NumberColumn("代次", min_value=1, step=1, width="small"),
                            "date": st.column_config.TextColumn("日期", width="small"),
                            "harvested_cells": st.column_config.NumberColumn("收获细胞数", min_value=0.0, format="%.2f", width="small"),
                            "inoculum_dishes": st.column_config.NumberColumn("种下盘数", min_value=1, step=1, width="small"),
                            "per_dish_cells": st.column_config.NumberColumn("每盘细胞数", min_value=0.0, format="%.2f", width="small"),
                            "status": st.column_config.SelectboxColumn("状态", options=["良好","一般","较差","污染"], width="small"),
                            "next_days": st.column_config.NumberColumn("间隔天数（仅最新记录）", min_value=0, step=1, width="small"),
                            "notes": st.column_config.TextColumn("备注", width="small")
                        },
                        use_container_width=True,
                        hide_index=True,
                        num_rows="fixed"
                    )

                    if st.button("批量保存修改", type="primary"):
                        errors = []
                        # 按批次分组更新，以便重新计算累计PD
                        for bid in edited_df["batch_name"].unique():
                            # 获取原始combined_df中属于该批次的记录ID，避免重复计算
                            batch_edited = edited_df[edited_df["batch_name"] == bid]
                            # 需要获取原始的batch_id，从combined_df中拿
                            original_batch_id = combined_df.loc[combined_df["batch_name"] == bid, "batch_id"].iloc[0]

                            # 逐行更新数据库
                            for _, row in batch_edited.iterrows():
                                rid = int(row["id"])
                                # 基本字段
                                passage = int(row["passage"])
                                date_str = str(row["date"])[:10]
                                try:
                                    date_obj = datetime.strptime(date_str, "%Y-%m-%d").date()
                                except:
                                    errors.append(f"ID {rid} 日期格式错误，请使用YYYY-MM-DD")
                                    continue
                                harvested = float(row["harvested_cells"])
                                dishes = int(row["inoculum_dishes"])
                                per_dish = float(row["per_dish_cells"])
                                inoculum = dishes * per_dish
                                status = row["status"]
                                next_days = int(row["next_days"])
                                notes = row["notes"] if row["notes"] else ""

                                # 更新记录
                                conn = get_connection()
                                # 判断该记录是否为该批次当前活跃记录（is_active=1）
                                active_rec = conn.execute("SELECT id FROM culture_records WHERE batch_id=? AND is_active=1", (original_batch_id,)).fetchone()
                                is_active = (active_rec is not None and active_rec[0] == rid)
                                next_date = None
                                if is_active and next_days > 0:
                                    next_date = (date_obj + timedelta(days=next_days)).isoformat()
                                conn.execute('''UPDATE culture_records
                                                SET passage=?, date=?, inoculum_cells=?, inoculum_dishes=?, harvested_cells=?,
                                                    status=?, next_passage_date=?, notes=?
                                                WHERE id=?''',
                                             (passage, date_str, inoculum, dishes, harvested, status, next_date, notes, rid))
                                conn.commit()
                                conn.close()
                            # 重算该批次的累计PD
                            recalculate_cumulative_pd(original_batch_id)
                        if errors:
                            for e in errors:
                                st.error(e)
                        st.success("所有修改已保存，累计PD已更新。")
                        st.rerun()

# ==================== 5. 添加冻存管（含孔位点选） ====================
elif menu == "❄️ 添加冻存管":
    st.header("添加冻存管记录")
    all_lines = get_all_cell_lines()
    if all_lines.empty:
        st.warning("请先添加细胞株。")
    else:
        # 细胞株选择在表单外
        line_name = st.selectbox("选择细胞株", all_lines["name"].tolist())
        selected_line_id = int(all_lines[all_lines["name"] == line_name]["id"].iloc[0])

        # ========= 新增：冻存盒孔位点选区 =========
                # ========= 多选孔位（9×9 网格点选版） =========
        with st.expander("📌 从冻存盒选择孔位（可多选）"):
            conn = get_connection()
            all_vials_df = pd.read_sql_query("SELECT location FROM frozen_vials", conn)
            conn.close()

            boxes = set()
            if not all_vials_df.empty:
                for loc in all_vials_df["location"]:
                    parts = loc.split("/")
                    if len(parts) >= 3:
                        boxes.add("/".join(parts[:3]))
            box_list = sorted(boxes)

            if box_list:
                selected_box = st.selectbox("选择已有冻存盒", box_list, key="box_selector")

                # 获取该盒子的占用情况
                conn = get_connection()
                box_vials = pd.read_sql_query('''
                    SELECT fv.id, fv.location, cl.name as cell_name, fv.status
                    FROM frozen_vials fv
                    JOIN cell_lines cl ON fv.cell_line_id = cl.id
                    WHERE fv.location LIKE ?
                ''', conn, params=(f"{selected_box}/%",))
                conn.close()

                # 计算占用孔位
                occupied_holes = {}
                for _, row in box_vials.iterrows():
                    hole = row["location"].split("/")[-1]
                    occupied_holes[hole] = {"cell_name": row["cell_name"], "status": row["status"]}

                # 9×9 孔位定义
                rows = "ABCDEFGHI"
                cols = list(range(1, 10))

                # 初始化或获取已选孔位（通过 session_state 管理多选集合）
                grid_key = f"selected_holes_{selected_box}"
                if grid_key not in st.session_state:
                    st.session_state[grid_key] = set()
                selected_set = st.session_state[grid_key]

                # 渲染网格，每个孔位是一个复选框
                st.caption(f"盒：{selected_box} | 已占用 {len(occupied_holes)} 孔，点击空格勾选")
                for r in rows:
                    cols_grid = st.columns(9)
                    for idx, c in enumerate(cols):
                        hole_id = f"{r}{c}"
                        with cols_grid[idx]:
                            if hole_id in occupied_holes:
                                vial = occupied_holes[hole_id]
                                label = f"🚫 {hole_id}"
                                st.checkbox(label, value=False, disabled=True,
                                            key=f"occ_{selected_box}_{hole_id}",
                                            help=f"{vial['cell_name']} ({vial['status']})")
                            else:
                                # 空闲孔位，可勾选
                                is_checked = hole_id in selected_set
                                cb = st.checkbox(f"⬜ {hole_id}", value=is_checked,
                                                 key=f"free_{selected_box}_{hole_id}")
                                if cb and hole_id not in selected_set:
                                    selected_set.add(hole_id)
                                elif not cb and hole_id in selected_set:
                                    selected_set.remove(hole_id)
                                # st.checkbox 的 on_change 不便，这里借助循环后的判断更新 session_state
                                # 注意：由于 Streamlit 重跑机制，这里 set 会即时更新，无需额外处理

                # 同步更新 session_state 中的选中集合
                st.session_state[grid_key] = selected_set

                # 生成选中的位置字符串（逗号分隔）
                if selected_set:
                    locations = [f"{selected_box}/{h}" for h in sorted(selected_set)]
                    st.session_state["selected_location"] = ",".join(locations)
                    st.success(f"已选中 {len(selected_set)} 个孔位：{', '.join(sorted(selected_set))}")
                else:
                    # 不选中任何孔时，清除之前可能残留的位置
                    if "selected_location" in st.session_state:
                        del st.session_state["selected_location"]

            else:
                st.info("暂无冻存盒，请先手动添加一支冻存管（位置格式：罐A/架1/盒1/A1）。")

        # ========= 原表单（位置自动填入并允许修改） =========
        with st.form("add_vial"):
            vial_count = st.number_input("添加管数", min_value=1, value=1, step=1)
            freeze_date = st.date_input("冻存日期", value=datetime.today())
            cell_count_per_vial = st.number_input("每管细胞数（×10⁵ cells）", min_value=0.0, format="%.2f")

            # 位置输入框，从点选结果自动填入，也可手动修改
            default_location = st.session_state.get("selected_location", "")
            location = st.text_input("冻存位置（格式：罐/架/盒/孔，例如：罐1/架1/盒1/A1）", value=default_location)

            pd_str = st.text_input("冻存时PD值（可留空，未知可不填）", value="")
            notes = st.text_area("备注")
            if st.form_submit_button("添加"):
                # ---- 解析 PD 值 ----
                pd_value = None
                if pd_str.strip():
                    try:
                        pd_value = float(pd_str)
                    except ValueError:
                        st.error("PD值格式错误，请输入数字或留空。")
                        st.stop()

                conn = get_connection()
                replace_id = st.session_state.get('replace_vial_id')
                if replace_id:
                    # 更新已有冻存管（已复苏孔位替换）
                    conn.execute('''UPDATE frozen_vials 
                                    SET cell_line_id = ?, freeze_date = ?, cell_count = ?, 
                                        location = ?, pd = ?, status = '在库', notes = ?
                                    WHERE id = ?''',
                                (selected_line_id, freeze_date.isoformat(), cell_count_per_vial, 
                                 location, pd_value, notes, replace_id))
                    st.success(f"已更新冻存管记录（孔位 {location}）")
                else:
                    # 插入新记录
                    for _ in range(vial_count):
                        conn.execute('''INSERT INTO frozen_vials
                            (cell_line_id, freeze_date, cell_count, location, pd, notes)
                            VALUES (?,?,?,?,?,?)''',
                            (selected_line_id, freeze_date.isoformat(), cell_count_per_vial, location, pd_value, notes))
                    st.success(f"已添加 {vial_count} 支冻存管。")
                conn.commit()
                conn.close()
                # 清除临时状态
                for key in ['selected_location', 'replace_vial_id']:
                    if key in st.session_state:
                        del st.session_state[key]
# ==================== 冻存盒视图 ====================
# ==================== 冻存盒管理（搜索编辑 / 网格视图） ====================
# ==================== 冻存管管理（网格 / 搜索编辑 / 导入Excel） ====================
elif menu == "🧊 冻存盒视图":
    st.header("冻存管管理")
    mode = st.radio("选择功能", ["📦 冻存盒网格视图", "🔍 搜索编辑", "📥 导入Excel"], horizontal=True)

    # ========================= 1. 网格视图 =========================
    if mode == "📦 冻存盒网格视图":
        conn = get_connection()
        df_all = pd.read_sql_query("""
            SELECT fv.id, fv.location, cl.name as cell_name, fv.freeze_date, fv.cell_count, fv.pd, fv.status
            FROM frozen_vials fv
            JOIN cell_lines cl ON fv.cell_line_id = cl.id
        """, conn)
        conn.close()

        if df_all.empty:
            st.info("暂无冻存管数据，请先添加或导入。")
        else:
            def extract_box_id(loc):
                parts = loc.split("/")
                if len(parts) >= 3:
                    return "/".join(parts[:3])
                return "格式不符"

            df_all["box_id"] = df_all["location"].apply(extract_box_id)
            df_valid = df_all[df_all["box_id"] != "格式不符"]
            if df_valid.empty:
                st.warning("没有符合“罐/架/盒/孔”格式的冻存管。")
            else:
                boxes = sorted(df_valid["box_id"].unique())
                selected_box = st.selectbox("选择冻存盒", boxes, key="box_view_selector")
                box_df = df_valid[df_valid["box_id"] == selected_box]

                st.subheader(f"{selected_box}  （共 {len(box_df)} 支管）")

                # 构建孔位映射
                hole_map = {}
                for _, row in box_df.iterrows():
                    hole = row["location"].split("/")[-1]
                    if hole.startswith("孔"):
                        hole = hole[1:]
                    hole_map[hole] = {"id": row["id"], "cell_name": row["cell_name"], "status": row["status"]}

                # 删除确认
                if "pending_delete" in st.session_state and st.session_state.pending_delete.get("box") == selected_box:
                    st.warning(f"确认删除孔位 {st.session_state.pending_delete['hole']} 的冻存管（{st.session_state.pending_delete['cell_name']}）吗？")
                    col1, col2 = st.columns(2)
                    if col1.button("✅ 确认删除", key="confirm_del_box"):
                        conn_del = get_connection()
                        conn_del.execute("DELETE FROM frozen_vials WHERE id = ?", (st.session_state.pending_delete["id"],))
                        conn_del.commit()
                        conn_del.close()
                        del st.session_state.pending_delete
                        st.success("已删除。")
                        st.rerun()
                    if col2.button("❌ 取消", key="cancel_del_box"):
                        del st.session_state.pending_delete
                        st.rerun()

                # 渲染 9x9 网格（HTML 表格）
                rows = "ABCDEFGHI"
                cols = list(range(1, 10))
                html = '<table style="border-collapse: collapse; width: 100%; font-size:12px;">'
                html += '<tr><th></th>' + ''.join([f'<th style="border:1px solid #ddd; background:#f0f0f0;">{c}</th>' for c in cols]) + '</tr>'
                for r in rows:
                    html += '<tr>'
                    html += f'<td style="font-weight:bold; border:1px solid #ddd; background:#f0f0f0; text-align:center;">{r}</td>'
                    for c in cols:
                        hole_id = f"{r}{c}"
                        if hole_id in hole_map:
                            vial = hole_map[hole_id]
                            status = vial['status']
                            bg = "#d4edda" if status == "在库" else "#fff3cd"
                            border = "#28a745" if status == "在库" else "#ffc107"
                            content = f"<b>{vial['cell_name']}</b><br>{status}"
                            html += f'<td style="border:2px solid {border}; background:{bg}; text-align:center; padding:2px;">{content}</td>'
                        else:
                            html += f'<td style="border:1px solid #ddd; color:#ccc; text-align:center;">＋</td>'
                    html += '</tr>'
                html += '</table>'
                st.markdown(html, unsafe_allow_html=True)
                st.markdown("🟢 在库  |  🟡 已复苏")

    # ========================= 2. 搜索编辑 =========================
    elif mode == "🔍 搜索编辑":
        st.subheader("搜索与编辑冻存管")
        search = st.text_input("输入细胞名称、盒子或备注关键词")

        if search.strip():
            like = f"%{search.strip()}%"
            conn = get_connection()
            df = pd.read_sql_query(
                "SELECT fv.id, fv.location, cl.name as cell_name, fv.cell_count, fv.pd, fv.freeze_date, fv.status, fv.notes "
                "FROM frozen_vials fv JOIN cell_lines cl ON fv.cell_line_id = cl.id "
                "WHERE cl.name LIKE ? OR fv.location LIKE ? OR fv.notes LIKE ? "
                "ORDER BY fv.location",
                conn, params=(like, like, like))
            conn.close()
        else:
            conn = get_connection()
            df = pd.read_sql_query(
                "SELECT fv.id, fv.location, cl.name as cell_name, fv.cell_count, fv.pd, fv.freeze_date, fv.status, fv.notes "
                "FROM frozen_vials fv JOIN cell_lines cl ON fv.cell_line_id = cl.id "
                "ORDER BY fv.location LIMIT 200", conn)
            conn.close()

        if df.empty:
            st.info("无匹配记录")
        else:
            st.write(f"找到 {len(df)} 条记录")
            # 表格展示
            display_df = df[['location', 'cell_name', 'cell_count', 'pd', 'freeze_date', 'status', 'notes']].copy()
            display_df.columns = ['位置', '细胞名', '细胞数', 'PD', '冻存日期', '状态', '原始信息']
            display_df['细胞数'] = display_df['细胞数'].apply(lambda x: f"{x:.2f}" if pd.notna(x) else '未知')
            display_df['PD'] = display_df['PD'].apply(lambda x: f"{x:.2f}" if pd.notna(x) else '无')
            display_df['原始信息'] = display_df['原始信息'].fillna('')
            st.dataframe(display_df, width='stretch', hide_index=True)

            st.markdown("---")
            # 选择记录进行编辑/删除
            options = [f"{row['location']} - {row['cell_name']} (ID:{row['id']})" for _, row in df.iterrows()]
            selected = st.selectbox("选择要操作的记录", options, key="select_record")
            if selected:
                selected_id = int(selected.split("(ID:")[-1].rstrip(')'))
                conn = get_connection()
                rec = conn.execute("SELECT * FROM frozen_vials WHERE id = ?", (selected_id,)).fetchone()
                conn.close()
                if rec:
                    col_names = ['id', 'cell_line_id', 'batch_id', 'freeze_date', 'cell_count', 'location', 'pd', 'status', 'notes']
                    vals = dict(zip(col_names, rec))
                    col1, col2 = st.columns(2)
                    with col1:
                        if st.button("✏️ 编辑", key=f"edit_vial_{selected_id}"):
                            st.session_state.edit_vial_id = selected_id
                            st.rerun()
                    with col2:
                        if st.button("🗑 删除", key=f"del_vial_{selected_id}"):
                            st.session_state.delete_vial_id = selected_id
                            st.rerun()

            # 编辑表单
            if 'edit_vial_id' in st.session_state:
                eid = st.session_state.edit_vial_id
                conn = get_connection()
                rec = conn.execute("SELECT * FROM frozen_vials WHERE id = ?", (eid,)).fetchone()
                conn.close()
                if rec:
                    col_names = ['id', 'cell_line_id', 'batch_id', 'freeze_date', 'cell_count', 'location', 'pd', 'status', 'notes']
                    vals = dict(zip(col_names, rec))
                    all_lines = get_all_cell_lines()
                    line_names = all_lines["name"].tolist() if not all_lines.empty else []
                    current_line_name = get_cell_line_info(vals['cell_line_id'])['name'] if vals['cell_line_id'] else ''

                    with st.form(key=f"edit_form_vial_{eid}"):
                        st.subheader(f"编辑 {vals['location']}")
                        if line_names:
                            default_idx = line_names.index(current_line_name) if current_line_name in line_names else 0
                            new_line = st.selectbox("细胞株", line_names, index=default_idx)
                        else:
                            new_line = st.text_input("细胞株（新建）", value=current_line_name)
                        new_location = st.text_input("位置", value=vals['location'])
                        new_count_str = st.text_input("细胞数（可留空）", value=str(vals['cell_count']) if vals['cell_count'] is not None else "")
                        new_pd = st.text_input("PD值（可留空）", value=str(vals['pd']) if vals['pd'] is not None else "")
                        new_freeze = st.date_input("冻存日期", value=datetime.strptime(vals['freeze_date'], '%Y-%m-%d') if vals['freeze_date'] else datetime.today())
                        new_status = st.selectbox("状态", ["在库", "已复苏", "已丢弃"], index=["在库", "已复苏", "已丢弃"].index(vals['status']) if vals['status'] in ["在库", "已复苏", "已丢弃"] else 0)
                        new_notes = st.text_area("备注", value=vals['notes'] if vals['notes'] else "")
                        if st.form_submit_button("保存"):
                            # 处理细胞株id
                            if line_names:
                                new_line_id = int(all_lines[all_lines["name"] == new_line]["id"].iloc[0])
                            else:
                                conn = get_connection()
                                try:
                                    conn.execute("INSERT INTO cell_lines (name, type) VALUES (?, '永生化')", (new_line,))
                                    conn.commit()
                                    new_line_id = conn.execute("SELECT last_insert_rowid()").fetchone()[0]
                                except:
                                    st.error("细胞株创建失败")
                                    st.stop()
                                finally:
                                    conn.close()
                            # 处理数字
                            new_count = None
                            if new_count_str.strip():
                                try:
                                    new_count = float(new_count_str)
                                except:
                                    st.error("细胞数格式错误")
                                    st.stop()
                            pd_num = None
                            if new_pd.strip():
                                try:
                                    pd_num = float(new_pd)
                                except:
                                    st.error("PD格式错误")
                                    st.stop()
                            conn_up = get_connection()
                            conn_up.execute('''UPDATE frozen_vials SET cell_line_id=?, location=?, cell_count=?, pd=?,
                                            freeze_date=?, status=?, notes=? WHERE id=?''',
                                         (new_line_id, new_location, new_count, pd_num,
                                          new_freeze.strftime('%Y-%m-%d'), new_status, new_notes, eid))
                            conn_up.commit()
                            conn_up.close()
                            del st.session_state.edit_vial_id
                            st.success("已保存")
                            st.rerun()
                        if st.button("取消编辑", key=f"cancel_edit_{eid}"):
                            del st.session_state.edit_vial_id
                            st.rerun()

            # 删除确认
            if 'delete_vial_id' in st.session_state:
                did = st.session_state.delete_vial_id
                conn = get_connection()
                rec = conn.execute("SELECT location, cell_line_id FROM frozen_vials WHERE id = ?", (did,)).fetchone()
                conn.close()
                if rec:
                    cell_name = get_cell_line_info(rec[1])["name"] if rec[1] else "?"
                    st.warning(f"确认删除 {rec[0]} 的 {cell_name}？")
                    c1, c2 = st.columns(2)
                    if c1.button("✅ 确认删除", key="confirm_del_vial"):
                        conn_del = get_connection()
                        conn_del.execute("DELETE FROM frozen_vials WHERE id = ?", (did,))
                        conn_del.commit()
                        conn_del.close()
                        del st.session_state.delete_vial_id
                        st.success("已删除")
                        st.rerun()
                    if c2.button("❌ 取消", key="cancel_del_vial"):
                        del st.session_state.delete_vial_id
                        st.rerun()

    # ========================= 3. 导入Excel =========================
    elif mode == "📥 导入Excel":
        st.subheader("从Excel导入冻存管")
        st.markdown("""
        **要求**：Excel 文件的每个 Sheet 代表一个冻存盒（Sheet 名即为盒子名称，如 `罐A/架1/盒1`），
        在 **A1:I9** 范围内为 9×9 孔位信息。每个格子的内容可多行，格式例如：

        细胞名称
        细胞数（如 1x10^7）
        PD值（可选）
        冻存日期
        复苏日期（如有，表示已复苏）

        支持 `PD=10.26`、`5.1X105`、`2022.4.15` 等常见写法。
        """)
        uploaded_file = st.file_uploader("上传 Excel 文件 (.xlsx)", type=["xlsx"])
        import_mode = st.radio("导入模式", ["追加（跳过已有孔位）", "覆盖（清空盒子后再导入）"], index=0)

        if uploaded_file:
            wb = openpyxl.load_workbook(uploaded_file, data_only=True)
            sheets = wb.sheetnames
            st.write(f"检测到 {len(sheets)} 个盒子：{', '.join(sheets)}")

            if st.button("开始导入"):
                total_inserted = 0
                total_skipped = 0
                errors = []
                for sheet_name in sheets:
                    ws = wb[sheet_name]
                    box = sheet_name.strip()
                    if not box:
                        continue
                    # 覆盖模式：先删除该盒子所有冻存管
                    if import_mode == "覆盖（清空盒子后再导入）":
                        conn = get_connection()
                        conn.execute("DELETE FROM frozen_vials WHERE location LIKE ?", (f"{box}/%",))
                        conn.commit()
                        conn.close()

                    for row_num in range(1, 10):
                        row_letter = chr(64 + row_num)  # A-I
                        for col_num in range(1, 10):
                            cell = ws.cell(row=row_num, column=col_num)
                            text = cell.value
                            if not text:
                                continue
                            parsed = parse_cell_text(text)
                            if not parsed:
                                continue
                            hole = f"{row_letter}{col_num}"
                            location = f"{box}/{hole}"

                            # 追加模式：检查是否已存在
                            if import_mode == "追加（跳过已有孔位）":
                                conn = get_connection()
                                exist = conn.execute("SELECT id FROM frozen_vials WHERE location = ?", (location,)).fetchone()
                                conn.close()
                                if exist:
                                    total_skipped += 1
                                    continue

                            # 获取或创建细胞株
                            cell_name = parsed['cell_name']
                            conn = get_connection()
                            line = conn.execute("SELECT id FROM cell_lines WHERE name = ?", (cell_name,)).fetchone()
                            if line:
                                line_id = line[0]
                            else:
                                # 自动创建细胞株（默认类型为“永生化”）
                                conn.execute("INSERT INTO cell_lines (name, type, revival_date) VALUES (?, '永生化', ?)",
                                             (cell_name, datetime.today().strftime('%Y-%m-%d')))
                                line_id = conn.execute("SELECT last_insert_rowid()").fetchone()[0]
                                conn.commit()
                            conn.close()

                            try:
                                conn = get_connection()
                                conn.execute('''INSERT INTO frozen_vials
                                    (cell_line_id, location, freeze_date, cell_count, pd, status, notes)
                                    VALUES (?,?,?,?,?,?,?)''',
                                    (line_id, location, parsed['freeze_date'], parsed['cell_count'],
                                     parsed['pd'], parsed['status'], parsed['notes']))
                                conn.commit()
                                conn.close()
                                total_inserted += 1
                            except Exception as e:
                                errors.append(f"{location}: {e}")
                st.success(f"导入完成：成功 {total_inserted} 条，跳过 {total_skipped} 条")
                if errors:
                    st.error("部分错误：")
                    for e in errors[:10]:
                        st.write(e)

# ==================== 6. 新建培养批次 ====================
elif menu == "🧪 新建培养批次":
    st.header("新建培养批次（复苏）")
    all_lines = get_all_cell_lines()
    if all_lines.empty:
        st.warning("请先添加细胞株。")
    else:
        line_name = st.selectbox("选择细胞株", all_lines["name"].tolist())
        selected_line_id = int(all_lines[all_lines["name"] == line_name]["id"].iloc[0])

        with st.form("add_batch"):
            revival_date = st.date_input("复苏日期", value=datetime.today())

            # 获取该细胞株所有在库冻存管
            vials_df = get_cell_line_frozen_vials(selected_line_id)
            in_storage = vials_df[vials_df['status'] == '在库'] if not vials_df.empty else pd.DataFrame()

            vial_options = ["不绑定（手动输入冻存日期）"]
            vial_id_map = {0: None}
            for i, (_, vial) in enumerate(in_storage.iterrows(), start=1):
                cell_str = f"{vial['cell_count']:.0f}" if vial['cell_count'] is not None else "?"
                label = f"{vial['freeze_date']} | {vial['location']} | {cell_str} cells"
                vial_options.append(label)
                vial_id_map[i] = vial['id']

            vial_choice = st.selectbox("选择来源冻存管（可选）", vial_options)
            selected_vial_id = vial_id_map[vial_options.index(vial_choice)]

            # 自动填入或手动输入冻存日期
            if selected_vial_id is not None:
                default_freeze_date = datetime.strptime(
                    in_storage.loc[in_storage['id'] == selected_vial_id, 'freeze_date'].values[0],
                    "%Y-%m-%d"
                ).date()
            else:
                default_freeze_date = datetime.today()
            source_freeze_date = st.date_input("所复苏冻存细胞的原始冻存日期", value=default_freeze_date)

            initial_pd = st.number_input("初始PD值", min_value=0.0, value=0.0, format="%.2f")
            initial_cell_count = st.number_input("复苏接种细胞数（可选）", min_value=0.0, value=0.0, format="%.2f")

            notes = st.text_area("批次备注（可选）")
            if st.form_submit_button("创建批次"):
                batch_name = f"{line_name}-{revival_date.strftime('%y%m%d')}"
                conn = get_connection()
                cursor = conn.cursor()
                cursor.execute('''INSERT INTO culture_batches
                    (cell_line_id, batch_name, revival_date, source_freeze_date,
                     initial_pd, initial_cell_count, is_active, notes)
                    VALUES (?,?,?,?,?,?,1,?)''',
                    (selected_line_id, batch_name, revival_date.isoformat(),
                     source_freeze_date.isoformat(), initial_pd,
                     initial_cell_count if initial_cell_count > 0 else None,
                     notes))
                # 获取新插入的批次 ID
                new_batch_id = cursor.lastrowid
                # 如果绑定了冻存管，更新其状态
                if selected_vial_id is not None:
                    cursor.execute(
                        "UPDATE frozen_vials SET status = '已复苏', batch_id = ? WHERE id = ?",
                        (new_batch_id, selected_vial_id)
                    )
                conn.commit()
                conn.close()
                st.success(f"已创建批次「{batch_name}」")
